#!/usr/bin/env python3
"""
Genera `Plantilla_Dashboard.xlsx`: una plantilla de Excel que NO necesita instalar nada
(ni Python ni macros). Se pegan los datos en la hoja "Datos" y el tablero se arma solo
con fórmulas:

  * Config  -> detecta automáticamente qué columna es la medida, la fecha, las
               categorías, la etapa (embudo), el departamento (mapa) y el costo
               (margen). Cada elección se puede cambiar con una lista desplegable.
  * Motor   -> cálculos intermedios (valores únicos, rankings, períodos...).
  * Dashboard -> KPIs, barras, líneas/área, matriz, dona, barras apiladas,
               mapa de Colombia, mapa de calor, dispersión y embudo.

Este script solo hace falta para REGENERAR la plantilla; quien la usa solo necesita Excel.

    python plantilla/generar_plantilla.py
"""
from __future__ import annotations

import sys
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
from dashboard_excel import (AREA_FILL, CARD, CRITICAL, FONT, GOOD, GRID, INK, INK_2, INK_3, MESES,  # noqa: E402
                             OTROS, SEQ_HI, SEQ_LO, SEQ_MID, SERIES, SURFACE, _font, _rich, estilo_chart)

SALIDA = Path(__file__).with_name("Plantilla_Dashboard.xlsx")
NCOLS = 40                                   # columnas de datos que se analizan (A:AN)
DR = "Datos!$A$2:$AN$1048576"                # rango de datos (sin encabezado)
HDR = "Datos!$A$1:$AN$1"
MAXU = 40                                    # valores únicos que se buscan por categoría
MAX_DISP = 5000                              # puntos de la dispersión
FMT_K = '[>=1000000]#,##0.0,,"M";[>=1000]#,##0.0,"K";#,##0'
FMT_N = "#,##0;-#,##0;;"                     # oculta los ceros en etiquetas
P0, P1 = 40, 40 + NCOLS - 1                  # filas del perfil de columnas en Config

HINTS = {
    "fecha": ["fecha", "date", "periodo", "hora", "time", "dia", "día"],
    "id": [" id ", " id", " cod", " nit ", " cedula", " cédula", " telefono", " teléfono", " celular", " sku",
           " factura", " numero", " número", " nro", " año ", " ano ", " year", " zip", " postal", " documento"],
    "med": ["venta", "ingreso", "monto", "importe", "total", "valor", "factur", "revenue", "sales", "amount",
            "utilidad", "ganancia", "gasto"],
    "pen": ["precio", "price", "unit", "costo", "cost", "promedio", "tasa", "%", "edad", "calific", "rating",
            "puntaje", "porcentaje", "satisf", "nps", "score", "nota"],
    "cat": ["categor", "region", "región", "producto", "linea", "línea", "segmento", "canal", "ciudad", "tienda",
            "sucursal", "vendedor", "cliente", "tipo", "marca", "zona", "sede"],
    "etapa": ["etapa", "fase", "estado", "stage", "embudo", "funnel", "paso", "status"],
    "depto": ["departamento", "depto", "dpto", "depart"],
    "costo": ["costo", "cost"],
}

# Mapa de mosaicos de Colombia: (nombre, abreviatura, patrón para SUMIFS, fila, columna)
DEPTOS = [
    ("San Andrés", "SAP", "*San Andr?s*", 1, 1), ("La Guajira", "LAG", "*Guajira", 1, 6),
    ("Atlántico", "ATL", "Atl?ntico", 2, 4), ("Magdalena", "MAG", "Magdalena", 2, 5), ("Cesar", "CES", "Cesar", 2, 6),
    ("Córdoba", "COR", "C?rdoba", 3, 3), ("Sucre", "SUC", "Sucre", 3, 4), ("Bolívar", "BOL", "Bol?var", 3, 5),
    ("Norte de Santander", "NSA", "Norte de Santander", 3, 6), ("Chocó", "CHO", "Choc?", 4, 2),
    ("Antioquia", "ANT", "Antioquia", 4, 3), ("Santander", "SAN", "Santander", 4, 5), ("Arauca", "ARA", "Arauca", 4, 6),
    ("Risaralda", "RIS", "Risaralda", 5, 2), ("Caldas", "CAL", "Caldas", 5, 3), ("Cundinamarca", "CUN", "Cundinamarca", 5, 4),
    ("Boyacá", "BOY", "Boyac?", 5, 5), ("Casanare", "CAS", "Casanare", 5, 6), ("Vichada", "VIC", "Vichada", 5, 7),
    ("Valle del Cauca", "VAC", "Valle*", 6, 2), ("Quindío", "QUI", "Quind?o", 6, 3), ("Tolima", "TOL", "Tolima", 6, 4),
    ("Bogotá D.C.", "BOG", "Bogot?*", 6, 5), ("Meta", "MET", "Meta", 6, 6), ("Guainía", "GUA", "Guain?a", 6, 7),
    ("Cauca", "CAU", "Cauca", 7, 2), ("Huila", "HUI", "Huila", 7, 3), ("Caquetá", "CAQ", "Caquet?", 7, 4),
    ("Guaviare", "GUV", "Guaviare", 7, 5), ("Vaupés", "VAU", "Vaup?s", 7, 6), ("Nariño", "NAR", "Nari?o", 8, 2),
    ("Putumayo", "PUT", "Putumayo", 8, 3), ("Amazonas", "AMA", "Amazonas", 8, 5),
]
DIAS = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]


def hint(celda: str, clave: str) -> str:
    arr = "{" + ",".join(f'"{w}"' for w in HINTS[clave]) + "}"
    return f"MIN(1,SUMPRODUCT(--ISNUMBER(SEARCH({arr},{celda}))))"


def mes(expr: str) -> str:
    return "CHOOSE(MONTH(" + expr + ")," + ",".join(f'"{m}"' for m in MESES) + ")"


def agg(crits: list[tuple[str, str]]) -> str:
    """Suma de la medida (o conteo de filas si no hay medida) con criterios."""
    if not crits:
        return "IF(cMed>0,SUM(rMed),nFilas)"
    c = ",".join(f"{r},{k}" for r, k in crits)
    return f"IF(cMed>0,SUMIFS(rMed,{c}),COUNTIFS({c}))"


def entre(c0: str, c1: str) -> list[tuple[str, str]]:
    return [("rFecha", f'">="&{c0}'), ("rFecha", f'"<"&{c1}')]


# ==========================================================================
class Plantilla:
    def __init__(self):
        self.wb = Workbook()
        self.dash = self.wb.active
        self.dash.title = "Dashboard"
        self.datos = self.wb.create_sheet("Datos")
        self.cfg = self.wb.create_sheet("Config")
        self.mot = self.wb.create_sheet("Motor")

    def nombre(self, n: str, ref: str):
        self.wb.defined_names[n] = DefinedName(n, attr_text=ref)

    # ------------------------------------------------------------------ Config
    def config(self):
        ws = self.cfg
        ws.sheet_properties.tabColor = "2A78D6"
        for col, w in {"A": 3, "B": 30, "C": 24, "D": 24, "E": 24, "F": 8}.items():
            ws.column_dimensions[col].width = w
        ws["B1"] = "Configuración del tablero"
        ws["B1"].font = _font(16, True)
        ws["B2"] = ("La plantilla elige sola las columnas (columna C). Si quieres otra, escógela en la columna "
                    "amarilla D; para volver a la automática, borra la celda amarilla.")
        ws["B2"].font = _font(9, color=INK_2)
        for j, h in enumerate(["Uso en el tablero", "Detección automática", "Tu elección (opcional)", "En uso", "Col."], 2):
            c = ws.cell(5, j, h)
            c.font = _font(9, True, "FFFFFF")
            c.fill = PatternFill("solid", fgColor="2A78D6")
        roles = [  # fila, etiqueta, nombre, columna de puntaje en el perfil
            (6, "Medida principal (lo que se suma)", "cMed", "I"),
            (7, "Medida secundaria (eje X dispersión)", "cMed2", "P"),
            (8, "Fecha", "cFecha", "L"),
            (9, "Categoría principal", "cCat1", "J"),
            (10, "Categoría secundaria", "cCat2", "Q"),
            (11, "Etapa (embudo)", "cEtapa", "M"),
            (12, "Departamento (mapa)", "cDepto", "N"),
            (13, "Costo (para el margen)", "cCosto", "R"),
        ]
        amarillo = PatternFill("solid", fgColor="FFF2B3")
        dv = DataValidation(type="list", formula1=f"=$B${P0}:$B${P1}", allow_blank=True)
        ws.add_data_validation(dv)
        for r, et, nm, sc in roles:
            ws.cell(r, 2, et).font = _font(10, True)
            rng = f"${sc}${P0}:${sc}${P1}"
            ws.cell(r, 3, f'=IF(MAX({rng})>0,INDEX($B${P0}:$B${P1},MATCH(MAX({rng}),{rng},0)),"")').font = _font(10, color=INK_2)
            d = ws.cell(r, 4)
            d.fill = amarillo
            d.font = _font(10)
            dv.add(d)
            ws.cell(r, 5, f'=IF(D{r}<>"",D{r},C{r})').font = _font(10, True)
            ws.cell(r, 6, f'=IF(E{r}="",0,IFERROR(MATCH(E{r},{HDR},0),0))').font = _font(9, color=INK_3)
            self.nombre(nm, f"Config!$F${r}")
        for r in range(6, 14):
            for j in range(2, 7):
                ws.cell(r, j).border = Border(bottom=Side(style="hair", color=GRID))

        escalares = [
            (16, "Filas de datos", f"=MAX(C{P0}:C{P1})", "nFilas", "#,##0"),
            (17, "Fecha mínima", '=IF(cFecha>0,MIN(rFecha),"")', "fMin", "dd/mm/yyyy"),
            (18, "Fecha máxima", '=IF(cFecha>0,MAX(rFecha),"")', "fMax", "dd/mm/yyyy"),
            (19, "Agrupación de la tendencia",
             '=IF(cFecha=0,"",IF(INT(fMax)-INT(fMin)<24,"Día",IF(INT(fMax)-INT(fMin)<168,"Semana",'
             'IF(INT(fMax)-INT(fMin)<730,"Mes",IF(INT(fMax)-INT(fMin)<2190,"Trimestre","Año")))))', "gran", None),
            (20, "La fecha incluye hora", "=IF(cFecha>0,IFERROR(SUMPRODUCT(--(MOD(rFecha,1)>0))>0,FALSE),FALSE)", "conHora", None),
            (21, "Total de la medida", "=" + agg([]), "total", "#,##0.##"),
            (22, "Valores distintos categoría principal", f"=SUMPRODUCT(--(Motor!$B$3:$B${2 + MAXU}<>\"\"))", "n1", None),
            (23, "Valores distintos categoría secundaria", f"=SUMPRODUCT(--(Motor!$K$3:$K${2 + MAXU}<>\"\"))", "n2", None),
            (24, "Dona usa la categoría secundaria", "=AND(cCat2>0,n2>=2,n2<n1)", "dona2", None),
            (25, "Inicio del último mes", '=IF(cFecha>0,DATE(YEAR(fMax),MONTH(fMax),1),"")', "mIni", "dd/mm/yyyy"),
            (26, "Nombre de la medida", '=IF(cMed>0,$E$6,"Registros")', "nomMed", None),
            (27, "Inicio del primer período",
             '=IF(cFecha=0,"",CHOOSE(gIdx,INT(fMin),INT(fMin)-WEEKDAY(fMin,2)+1,DATE(YEAR(fMin),MONTH(fMin),1),'
             'DATE(YEAR(fMin),3*INT((MONTH(fMin)-1)/3)+1,1),DATE(YEAR(fMin),1,1)))', "tBase", "dd/mm/yyyy"),
            (28, "Índice de agrupación", '=IFERROR(MATCH(gran,{"Día","Semana","Mes","Trimestre","Año"},0),3)', "gIdx", None),
        ]
        ws["B15"] = "Valores calculados"
        ws["B15"].font = _font(11, True)
        for r, et, f, nm, fmt in escalares:
            ws.cell(r, 2, et).font = _font(10)
            c = ws.cell(r, 3, f)
            c.font = _font(10, color=INK_2)
            if fmt:
                c.number_format = fmt
            self.nombre(nm, f"Config!$C${r}")

        # Rangos dinámicos: de la fila 2 a la última fila con datos, en la columna elegida
        for nm, c in [("rMed", "cMed"), ("rMed2", "cMed2"), ("rFecha", "cFecha"), ("rCat1", "cCat1"),
                      ("rCat2", "cCat2"), ("rEtapa", "cEtapa"), ("rDepto", "cDepto"), ("rCosto", "cCosto")]:
            self.nombre(nm, f"INDEX({DR},1,MAX(1,{c})):INDEX({DR},MAX(1,nFilas),MAX(1,{c}))")

        # Perfil de cada columna de Datos
        ws.cell(P0 - 2, 2, "Perfil automático de columnas (no editar)").font = _font(11, True)
        enc = ["#", "Encabezado", "Celdas", "Numéricas", "¿Fecha?", "¿Número?", "Texto norm.", "¿Texto?",
               "P. medida", "P. categoría", "Distintos (500)", "P. fecha", "P. etapa", "P. depto", "P. costo",
               "P. medida 2", "P. categoría 2", "P. costo 2"]
        for j, h in enumerate(enc, 1):
            c = ws.cell(P0 - 1, j, h)
            c.font = _font(8, True, INK_2)
        for j in range(1, NCOLS + 1):
            r = P0 + j - 1
            peso = f"+{(NCOLS + 1 - j) / 1000:.3f}"
            f = {
                "A": j,
                "B": f'=IF(INDEX({HDR},1,{j})="","",INDEX({HDR},1,{j})&"")',
                "C": f'=IF(B{r}="",0,COUNTA(INDEX({DR},0,{j})))',
                "D": f'=IF(B{r}="",0,COUNT(INDEX({DR},0,{j})))',
                "G": f'=" "&LOWER(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(B{r},"_"," "),"-"," "),"."," "))&" "',
                "E": (f'=AND(C{r}>0,D{r}>=0.9*C{r},OR(LEFT(CELL("format",INDEX(Datos!$A$2:$AN$2,1,{j})),1)="D",'
                      f'{hint(f"G{r}", "fecha")}>0))'),
                "F": f'=AND(C{r}>0,D{r}>=0.9*C{r},NOT(E{r}),{hint(f"G{r}", "id")}=0)',
                "H": f"=AND(C{r}>0,D{r}<0.5*C{r})",
                "K": (f'=IF(H{r},SUMPRODUCT((INDEX(Datos!$A$2:$AN$501,0,{j})<>"")/'
                      f'COUNTIF(INDEX(Datos!$A$2:$AN$501,0,{j}),INDEX(Datos!$A$2:$AN$501,0,{j})&"")),0)'),
                "I": f'=IF(F{r},2+10*{hint(f"G{r}", "med")}-{hint(f"G{r}", "pen")}{peso},0)',
                "J": f'=IF(AND(H{r},K{r}>=2),IF(K{r}<=40,2+10*{hint(f"G{r}", "cat")},0.5+{hint(f"G{r}", "cat")}){peso},0)',
                "L": f'=IF(E{r},1+{hint(f"G{r}", "fecha")}{peso},0)',
                "M": f'=IF(AND(H{r},K{r}>=2,K{r}<=15),{hint(f"G{r}", "etapa")}*(1{peso}),0)',
                "N": f'=IF(H{r},{hint(f"G{r}", "depto")}*(1{peso}),0)',
                "O": f'=IF(F{r},{hint(f"G{r}", "costo")}*(1{peso}),0)',
                "P": f'=IF(B{r}=$E$6,0,I{r})',
                "Q": f'=IF(B{r}=$E$9,0,J{r})',
                "R": f'=IF(OR(B{r}=$E$6,INDEX($I${P0}:$I${P1},MAX(1,$F$6))<10),0,O{r})',
            }
            for col, v in f.items():
                c = ws[f"{col}{r}"]
                c.value = v
                c.font = _font(8, color=INK_3)
                if col == "K":
                    c.number_format = "0"
                if col in "IJLMNOPQR":
                    c.number_format = "0.000"
        for col in "GHIJKLMNOPQR":
            ws.column_dimensions[col].width = 10
        ws.freeze_panes = "A6"

    # ------------------------------------------------------------------ Motor
    def bloque_unicos(self, c0: int, rango: str, flag: str, titulo: str, n: int = MAXU, conteo: bool = False):
        """Columnas c0..c0+7: únicos, valor, rango, k, etiqueta ordenada, valor ordenado, % del total."""
        ws = self.mot
        A, B, C, D, E, F, G, H = (L(c0 + i) for i in range(8))
        ws[f"{A}1"] = titulo
        ws[f"{A}1"].font = _font(10, True)
        s, e = 3, 2 + n
        for i in range(n):
            r = s + i
            ws[f"{A}{r}"] = i + 1
            ws[f"{B}{r}"] = ArrayFormula(
                f"{B}{r}",
                f'=IF({flag}=0,"",IFERROR(INDEX({rango},MATCH(0,COUNTIF({B}$2:{B}{r - 1},{rango}&"")+({rango}&""=""),0))&"",""))')
            val = f"COUNTIFS({rango},{B}{r})" if conteo else agg([(rango, f"{B}{r}")])
            ws[f"{C}{r}"] = f'=IF({B}{r}="","",{val})'
            ws[f"{D}{r}"] = f'=IF({B}{r}="","",COUNTIF({C}${s}:{C}${e},">"&{C}{r})+COUNTIF({C}${s}:{C}{r},{C}{r}))'
            ws[f"{E}{r}"] = i + 1
            ws[f"{F}{r}"] = f'=IFERROR(INDEX({B}${s}:{B}${e},MATCH({E}{r},{D}${s}:{D}${e},0)),"")'
            ws[f"{G}{r}"] = f'=IFERROR(INDEX({C}${s}:{C}${e},MATCH({E}{r},{D}${s}:{D}${e},0)),"")'
            tot = f"SUM({C}${s}:{C}${e})" if conteo else "total"
            ws[f"{H}{r}"] = f'=IF({G}{r}="","",IFERROR({G}{r}/{tot},""))'
            ws[f"{G}{r}"].number_format = ws[f"{C}{r}"].number_format = "#,##0.##"
            ws[f"{H}{r}"].number_format = "0.0%"
        for j, h in enumerate(["k", "Único", "Valor", "Rango", "k", "Ordenado", "Valor", "%"]):
            ws.cell(2, c0 + j).value = None
        return {"lab": F, "val": G, "pct": H, "s": s}

    def motor(self):
        ws = self.mot
        ws.sheet_properties.tabColor = "B4B2AC"
        ws["A1"].font = _font(10, True)
        self.u1 = self.bloque_unicos(1, "rCat1", "cCat1", "Categoría principal")      # A..H
        self.u2 = self.bloque_unicos(10, "rCat2", "cCat2", "Categoría secundaria")    # J..Q
        self.ue = self.bloque_unicos(19, "rEtapa", "cEtapa", "Etapas", n=15, conteo=True)  # S..Z

        # Tendencia: hasta 24 períodos desde el inicio de los datos; la agrupación se adapta  AB..AF
        ws["AB1"] = "Tendencia"
        ws["AB1"].font = _font(10, True)
        for k in range(1, 25):
            r = k + 2
            j = k - 1
            ini = (f"CHOOSE(gIdx,tBase+{j},tBase+{7 * j},DATE(YEAR(tBase),MONTH(tBase)+{j},1),"
                   f"DATE(YEAR(tBase),MONTH(tBase)+{3 * j},1),DATE(YEAR(tBase)+{j},1,1))")
            ws[f"AB{r}"] = f'=IF(cFecha=0,"",IF({ini}>fMax,"",{ini}))'
            ws[f"AC{r}"] = (f'=IF(AB{r}="","",CHOOSE(gIdx,AB{r}+1,AB{r}+7,DATE(YEAR(AB{r}),MONTH(AB{r})+1,1),'
                            f'DATE(YEAR(AB{r}),MONTH(AB{r})+3,1),DATE(YEAR(AB{r})+1,1,1)))')
            ws[f"AD{r}"] = f'=IF(AB{r}="","",{mes(f"AB{r}")}&" "&RIGHT(YEAR(AB{r}),2))'
            ws[f"AE{r}"] = f'=IF(AB{r}="",NA(),{agg(entre(f"AB{r}", f"AC{r}"))})'
            ws[f"AF{r}"] = f'=IF(AB{r}="",NA(),AB{r})'
            ws[f"AB{r}"].number_format = ws[f"AC{r}"].number_format = "dd/mm/yyyy"
        ws["AE2"] = "=nomMed"

        # Dona: top 5 + Otros  AG..AH
        ws["AG1"] = "Participación"
        ws["AG1"].font = _font(10, True)
        u1, u2 = self.u1, self.u2
        for k in range(1, 6):
            r = k + 2
            ws[f"AG{r}"] = f"=IF(dona2,{u2['lab']}{r},{u1['lab']}{r})"
            ws[f"AH{r}"] = f'=IF(AG{r}="","",IF(dona2,{u2["val"]}{r},{u1["val"]}{r}))'
        ws["AG8"] = '=IF(IF(dona2,n2,n1)>5,"Otros","")'
        ws["AH8"] = '=IF(AG8="","",MAX(0,total-SUM(AH3:AH7)))'
        ws["AH2"] = "=nomMed"

        # Barras apiladas: 12 meses x top 4 (+ Otros)  AJ..AR
        ws["AJ1"] = "Composición mensual"
        ws["AJ1"].font = _font(10, True)
        series = ["AM", "AN", "AO", "AP"]
        for k, col in enumerate(series, 1):
            ws[f"{col}2"] = f"=IF(cCat2>0,{u2['lab']}{k + 2},{u1['lab']}{k + 2})"
        ws["AQ2"] = "Otros"
        ws["AR2"] = "Total"
        for k in range(1, 13):
            r = k + 2
            ws[f"AJ{r}"] = f'=IF(cFecha=0,"",DATE(YEAR(fMax),MONTH(fMax)-12+{k},1))'
            ws[f"AK{r}"] = f'=IF(AJ{r}="","",DATE(YEAR(AJ{r}),MONTH(AJ{r})+1,1))'
            ws[f"AL{r}"] = f'=IF(AJ{r}="","",{mes(f"AJ{r}")}&" "&RIGHT(YEAR(AJ{r}),2))'
            b = entre(f"$AJ{r}", f"$AK{r}")
            for col in series:
                a2 = agg(b + [("rCat2", f"{col}$2")])
                a1 = agg(b + [("rCat1", f"{col}$2")])
                ws[f"{col}{r}"] = f'=IF(OR($AJ{r}="",{col}$2=""),0,IF(cCat2>0,{a2},{a1}))'
            ws[f"AR{r}"] = f'=IF(AJ{r}="",0,{agg(b)})'
            ws[f"AQ{r}"] = f"=MAX(0,AR{r}-SUM(AM{r}:AP{r}))"

        # Matriz: top 10 categoría principal x últimos 6 meses + total  AT..BA
        ws["AT1"] = "Matriz"
        ws["AT1"].font = _font(10, True)
        meses_cols = ["AU", "AV", "AW", "AX", "AY", "AZ"]
        for k, col in enumerate(meses_cols, 1):
            ws[f"{col}1"] = f'=IF(cFecha=0,"",DATE(YEAR(fMax),MONTH(fMax)-6+{k},1))'
            ws[f"{col}1"].number_format = "dd/mm/yyyy"
            ws[f"{col}2"] = f'=IF({col}1="","",{mes(f"{col}1")}&" "&RIGHT(YEAR({col}1),2))'
        ws["AT2"] = '=IF(cCat1>0,Config!$E$9,"")'
        ws["BA2"] = '="Total 6 meses"'
        for i in range(10):
            r = i + 3
            ws[f"AT{r}"] = f"={u1['lab']}{r}"
            for col in meses_cols:
                crit = [("rCat1", f"$AT{r}")] + entre(f"{col}$1", f"DATE(YEAR({col}$1),MONTH({col}$1)+1,1)")
                ws[f"{col}{r}"] = f'=IF(OR($AT{r}="",{col}$1=""),"",{agg(crit)})'
            ws[f"BA{r}"] = f'=IF(AT{r}="","",SUM(AU{r}:AZ{r}))'

        # Mapa de calor: horas (o meses) x día de la semana  BC..BJ
        ws["BC1"] = "Mapa de calor"
        ws["BC1"].font = _font(10, True)
        dias_cols = [L(56 + d) for d in range(7)]   # BD..BJ
        for d, col in enumerate(dias_cols):
            ws[f"{col}2"] = DIAS[d]
        ws["BC2"] = '=IF(conHora,"Hora","Mes")'
        vals = "IF(cMed>0,IF(ISNUMBER(rMed),rMed,0),1)"
        for k in range(1, 25):
            r = k + 2
            ws[f"BC{r}"] = (f'=IF(conHora,RIGHT("0"&{k - 1},2)&":00",IF({k}<=12,'
                            f'CHOOSE({min(k, 12)},{",".join(chr(34) + m + chr(34) for m in MESES)}),""))')
            for d, col in enumerate(dias_cols, 1):
                f = (f'=IF(OR(cFecha=0,AND(NOT(conHora),{k}>12)),"",SUMPRODUCT(IFERROR(ISNUMBER(rFecha)*'
                     f'(WEEKDAY(rFecha,2)={d})*IF(conHora,HOUR(rFecha)={k - 1},MONTH(rFecha)={k})*{vals},0)))')
                ws[f"{col}{r}"] = ArrayFormula(f"{col}{r}", f)

        # Mapa de Colombia  BL..BP
        ws["BL1"] = "Mapa por departamento"
        ws["BL1"].font = _font(10, True)
        n = len(DEPTOS)
        for i, (nom, ab, pat, _, _) in enumerate(DEPTOS):
            r = i + 3
            ws[f"BL{r}"] = nom
            ws[f"BM{r}"] = pat
            ws[f"BN{r}"] = f'=IF(cDepto=0,"",{agg([("rDepto", f"BM{r}")])})'
            ws[f"BO{r}"] = (f'=IF(BN{r}="","",COUNTIF(BN$3:BN${n + 2},">"&BN{r})+COUNTIF(BN$3:BN{r},BN{r}))')
        for k in range(1, 6):
            r = k + 2
            ws[f"BP{r}"] = f'=IFERROR(INDEX(BL$3:BL${n + 2},MATCH({k},BO$3:BO${n + 2},0)),"")'
            ws[f"BQ{r}"] = f'=IFERROR(INDEX(BN$3:BN${n + 2},MATCH({k},BO$3:BO${n + 2},0)),"")'
        self.n_deptos = n

        # Embudo: etapas ordenadas por cantidad  BS..BV
        ws["BS1"] = "Embudo"
        ws["BS1"].font = _font(10, True)
        ue = self.ue
        ws["BU2"] = "Registros"
        ws["BT2"] = "Relleno"
        for k in range(1, 7):
            r = k + 2
            ws[f"BS{r}"] = f'=IF(cEtapa=0,"",{ue["lab"]}{r})'
            ws[f"BU{r}"] = f'=IF(BS{r}="","",{ue["val"]}{r})'
            ws[f"BT{r}"] = f'=IF(BU{r}="","",(MAX(BU$3:BU$8)-BU{r})/2)'
            ws[f"BV{r}"] = f'=IF(BU{r}="","",BU{r}/BU$3)'
            ws[f"BV{r}"].number_format = "0.0%"

        # Dispersión  BX..BY
        ws["BX1"] = "Dispersión"
        ws["BX1"].font = _font(10, True)
        for k in range(1, MAX_DISP + 1):
            r = k + 2
            ok = f"AND(ISNUMBER(INDEX(rMed2,{k})),ISNUMBER(INDEX(rMed,{k})))"
            ws[f"BX{r}"] = f"=IF(OR(cMed2=0,{k}>nFilas),NA(),IF({ok},INDEX(rMed2,{k}),NA()))"
            ws[f"BY{r}"] = f"=IF(OR(cMed2=0,{k}>nFilas),NA(),IF({ok},INDEX(rMed,{k}),NA()))"

        # KPIs  CA..CB
        ws["CA1"] = "KPIs"
        ws["CA1"].font = _font(10, True)
        ini, fin, ant = "mIni", "DATE(YEAR(mIni),MONTH(mIni)+1,1)", "DATE(YEAR(mIni),MONTH(mIni)-1,1)"
        kp = [
            ("Actual", f'=IF(cFecha=0,0,{agg(entre(ini, fin))})'),
            ("Anterior", f'=IF(cFecha=0,0,{agg(entre(ant, ini))})'),
            ("Filas actual", f'=IF(cFecha=0,0,COUNTIFS({",".join(f"{a},{b}" for a, b in entre(ini, fin))}))'),
            ("Filas anterior", f'=IF(cFecha=0,0,COUNTIFS({",".join(f"{a},{b}" for a, b in entre(ant, ini))}))'),
            ("Costo actual", f'=IF(OR(cFecha=0,cCosto=0),0,SUMIFS(rCosto,{",".join(f"{a},{b}" for a, b in entre(ini, fin))}))'),
            ("Costo anterior", f'=IF(OR(cFecha=0,cCosto=0),0,SUMIFS(rCosto,{",".join(f"{a},{b}" for a, b in entre(ant, ini))}))'),
            ("Mes anterior", f'=IF(cFecha=0,"",{mes(ini)}&" "&RIGHT(YEAR({ini}),2)&" vs "&{mes(ant)}&" "&RIGHT(YEAR({ant}),2))'),
            ("Margen total", '=IF(OR(cCosto=0,cMed=0),"",IFERROR((SUM(rMed)-SUM(rCosto))/SUM(rMed),""))'),
        ]
        for i, (et, f) in enumerate(kp):
            ws.cell(i + 3, 79, et)
            ws.cell(i + 3, 80, f)
        self.kpi = {et: f"Motor!$CB${i + 3}" for i, (et, _) in enumerate(kp)}
        ws.column_dimensions["A"].width = 6

    # ------------------------------------------------------------------ Datos de ejemplo
    def datos_ejemplo(self, n=2400, semilla=11):
        rng = np.random.default_rng(semilla)
        ws = self.datos
        deptos = {"Antioquia": "Medellín", "Bogotá D.C.": "Bogotá", "Valle del Cauca": "Cali",
                  "Atlántico": "Barranquilla", "Santander": "Bucaramanga", "Bolívar": "Cartagena",
                  "Risaralda": "Pereira", "Nariño": "Pasto", "Meta": "Villavicencio", "Huila": "Neiva"}
        pesos = np.array([.2, .26, .14, .1, .07, .07, .05, .04, .04, .03])
        cats = ["Tecnología", "Hogar", "Moda", "Deportes", "Belleza", "Juguetes"]
        canales = ["Web", "App", "Tienda física", "Marketplace"]
        etapas = ["1. Visita", "2. Carrito", "3. Pago", "4. Compra"]
        precios = {"Tecnología": 850000, "Hogar": 210000, "Moda": 120000, "Deportes": 160000,
                   "Belleza": 70000, "Juguetes": 90000}
        enc = ["ID Pedido", "Fecha", "Departamento", "Ciudad", "Categoría", "Canal", "Etapa",
               "Unidades", "Precio unitario", "Ventas", "Costo"]
        ws.append(enc)
        inicio = datetime(2024, 4, 1)
        filas = []
        for i in range(n):
            dia = int(rng.integers(0, 540))
            hora = int(rng.choice(range(24), p=_horas()))
            f = inicio + timedelta(days=dia, hours=hora, minutes=int(rng.integers(0, 60)))
            dep = rng.choice(list(deptos), p=pesos)
            cat = rng.choice(cats)
            u = int(rng.integers(1, 6))
            p = round(precios[cat] * rng.uniform(.7, 1.3), -2)
            temporada = 1 + .4 * (f.month in (11, 12)) + .15 * (f.weekday() >= 4)
            v = round(u * p * temporada)
            filas.append([10001 + i, f, dep, deptos[dep], cat, rng.choice(canales, p=[.35, .3, .25, .1]),
                          rng.choice(etapas, p=[.45, .27, .17, .11]), u, p, v, round(v * rng.uniform(.55, .75))])
        filas.sort(key=lambda x: x[1])
        for x in filas:
            ws.append(x)
        for (c,) in ws.iter_rows(min_row=2, max_row=n + 1, min_col=2, max_col=2):
            c.number_format = "dd/mm/yyyy hh:mm"
        for col in ("I", "J", "K"):
            for (c,) in ws.iter_rows(min_row=2, max_row=n + 1, min_col=ord(col) - 64, max_col=ord(col) - 64):
                c.number_format = "#,##0"
        for j, h in enumerate(enc, 1):
            ws.cell(1, j).font = _font(10, True, "FFFFFF")
            ws.cell(1, j).fill = PatternFill("solid", fgColor="2A78D6")
            ws.column_dimensions[L(j)].width = max(12, len(h) + 4)
        ws.column_dimensions["B"].width = 17
        ws.column_dimensions["C"].width = 18
        ws.freeze_panes = "A2"

    # ------------------------------------------------------------------ Dashboard
    def dashboard(self):
        ws = self.dash
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90
        # Rejilla: A margen | B..H panel izquierdo | I separador | J..P panel derecho | Q margen
        for j in range(1, 18):
            ws.column_dimensions[L(j)].width = 2.5 if L(j) in ("A", "I", "Q") else 10.5
        self.izq = list(range(2, 9))
        self.der = list(range(10, 17))
        self.todo = self.izq + self.der
        fondo = PatternFill("solid", fgColor=SURFACE)
        for fila in ws.iter_rows(min_row=1, max_row=190, min_col=1, max_col=17):
            for c in fila:
                c.fill = fondo
        ws.row_dimensions[2].height = 30
        ws["B2"] = '="Dashboard de "&nomMed'
        ws["B2"].font = _font(20, True)
        ws["B3"] = ('="Pega tus datos en la hoja Datos (encabezados en la fila 1) · "&TEXT(nFilas,"#,##0")&'
                    '" registros · columnas elegidas en la hoja Config"')
        ws["B3"].font = _font(9, color=INK_2)
        self.kpis()
        y = 11
        alto_chart = 16
        y = self.fila_paneles(y, [self.p_tendencia], alto_chart, completo=True)
        y = self.fila_paneles(y, [self.p_ranking, self.p_dona], alto_chart)
        y = self.fila_paneles(y, [self.p_apiladas, self.p_embudo], alto_chart)
        y = self.fila_paneles(y, [self.p_matriz], 12, completo=True)
        y = self.fila_paneles(y, [self.p_mapa, self.p_top_deptos], 17)
        y = self.fila_paneles(y, [self.p_dispersion, self.p_ranking2], alto_chart)
        y = self.fila_paneles(y, [self.p_calor], 26, completo=True)
        ws.print_area = f"A1:Q{y}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    def tarjeta(self, cols, f0, f1):
        ws = self.dash
        blanco = PatternFill("solid", fgColor=CARD)
        borde = Side(style="thin", color=GRID)
        for r in range(f0, f1 + 1):
            for c in cols:
                cel = ws.cell(r, c)
                cel.fill = blanco
                cel.border = Border(top=borde if r == f0 else None, bottom=borde if r == f1 else None,
                                    left=borde if c == cols[0] else None, right=borde if c == cols[-1] else None)

    def kpis(self):
        ws = self.dash
        k = self.kpi
        var = lambda a, b, txt: (f'=IF(OR(cFecha=0,N({b})=0),"",IF({a}>={b},"▲ ","▼ ")&'
                                 f'ROUND(ABS({a}/{b}-1)*100,1)&"% · "&{txt})')
        prom_a = f'IFERROR({k["Actual"]}/{k["Filas actual"]},0)'
        prom_b = f'IFERROR({k["Anterior"]}/{k["Filas anterior"]},0)'
        marg_a = f'IFERROR(({k["Actual"]}-{k["Costo actual"]})/{k["Actual"]},0)'
        marg_b = f'IFERROR(({k["Anterior"]}-{k["Costo anterior"]})/{k["Anterior"]},0)'
        tarjetas = [
            ('=UPPER(nomMed)&" TOTAL"', "=total", FMT_K, var(k["Actual"], k["Anterior"], k["Mes anterior"])),
            ('="REGISTROS"', "=nFilas", "#,##0", var(k["Filas actual"], k["Filas anterior"], k["Mes anterior"])),
            ('="PROMEDIO POR REGISTRO"', "=IFERROR(total/nFilas,0)", FMT_K, var(prom_a, prom_b, k["Mes anterior"])),
            ('=IF(cCosto>0,"MARGEN",IF(cCat1>0,UPPER($E$9)&" DISTINTOS","—"))'.replace("$E$9", "Config!$E$9"),
             f'=IF(cCosto>0,ROUND({k["Margen total"]}*100,1)&"%",IF(cCat1>0,IF(n1>={MAXU},n1&"+",n1&""),""))',
             "@",
             (f'=IF(OR(cCosto=0,cFecha=0,N({k["Anterior"]})=0),"",IF({marg_a}>={marg_b},"▲ ","▼ ")&'
              f'ROUND(ABS({marg_a}-{marg_b})*100,1)&" pp · "&{k["Mes anterior"]})')),
        ]
        slots = [(2, 4), (6, 8), (10, 12), (14, 16)]
        for (a, b), (t, v, fmt, d) in zip(slots, tarjetas):
            cols = list(range(a, b + 1))
            self.tarjeta(cols, 5, 8)
            ws.merge_cells(start_row=5, start_column=a, end_row=5, end_column=b)
            ws.merge_cells(start_row=6, start_column=a, end_row=7, end_column=b)
            ws.merge_cells(start_row=8, start_column=a, end_row=8, end_column=b)
            c = ws.cell(5, a, t)
            c.font = _font(8, True, INK_2)
            c.alignment = Alignment(indent=1, vertical="bottom", shrink_to_fit=True)
            c = ws.cell(6, a, v)
            c.font = _font(22, True)
            c.number_format = fmt
            c.alignment = Alignment(horizontal="left", indent=1, vertical="center", shrink_to_fit=True)
            c = ws.cell(8, a, d)
            c.font = _font(9, color=INK_2)
            c.alignment = Alignment(indent=1, vertical="top", shrink_to_fit=True)
            ref = f"{L(a)}8"
            ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEFT({ref},1)="▲"'], font=Font(name=FONT, size=9, bold=True, color=GOOD)))
            ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEFT({ref},1)="▼"'], font=Font(name=FONT, size=9, bold=True, color=CRITICAL)))
        for r, h in ((5, 18), (6, 22), (7, 22), (8, 18)):
            ws.row_dimensions[r].height = h

    def fila_paneles(self, y, paneles, alto, completo=False):
        ws = self.dash
        grupos = [self.todo] if completo else [self.izq, self.der]
        for cols, fn in zip(grupos, paneles):
            self.tarjeta(list(range(cols[0], cols[-1] + 1)), y, y + alto + 2)
            titulo, sub = fn(cols, y + 2, alto, completo)
            c = ws.cell(y, cols[0], titulo)
            c.font = _font(12, True)
            c.alignment = Alignment(indent=1, vertical="bottom")
            c = ws.cell(y + 1, cols[0], sub)
            c.font = _font(9, color=INK_2)
            c.alignment = Alignment(indent=1)
        ws.row_dimensions[y].height = 22
        return y + alto + 4

    def poner(self, ch, cols, fila, alto, completo):
        px = 10.5 * 7 + 5
        ancho = (14 * px + 19) if completo else 7 * px
        ch.width = (ancho - 14) * 0.026458
        ch.height = alto * 20 * 0.026458
        self.dash.add_chart(ch, f"{L(cols[0])}{fila}")

    @staticmethod
    def aviso(cond: str, msg: str, normal: str) -> str:
        return f'=IF({cond},"⚠ {msg}",{normal})'

    # --- paneles
    def p_tendencia(self, cols, fila, alto, completo):
        # Línea XY sobre fechas reales: el eje se ajusta solo al rango de los datos y omite los vacíos
        ch = ScatterChart()
        estilo_chart(ch)
        ch.scatterStyle = "lineMarker"
        s = Series(Reference(self.mot, min_col=31, min_row=3, max_row=26),        # AE
                   Reference(self.mot, min_col=32, min_row=3, max_row=26), title="Tendencia")   # AF
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=SERIES[0], w=31750))
        s.marker = Marker(symbol="circle", size=6)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=SERIES[0])
        s.marker.graphicalProperties.line.solidFill = CARD
        s.smooth = False
        ch.series.append(s)
        ch.x_axis.numFmt = "mmm yy"
        ch.y_axis.numFmt = FMT_K
        ch.y_axis.scaling.min = 0
        self.poner(ch, cols, fila, alto, completo)
        return ('="Tendencia de "&nomMed',
                self.aviso("cFecha=0", "No hay columna de fecha: elígela en Config",
                           '"Total por "&LOWER(gran)&" · se agrupa solo según el rango de fechas"'))

    def _barras_h(self, blk, color):
        ch = BarChart()
        ch.type = "bar"
        estilo_chart(ch, grid_y=False)
        c_val = self._col(blk["val"])
        c_lab = self._col(blk["lab"])
        ch.add_data(Reference(self.mot, min_col=c_val, min_row=3, max_row=12), titles_from_data=False)
        ch.set_categories(Reference(self.mot, min_col=c_lab, min_row=3, max_row=12))
        ch.gapWidth = 45
        s = ch.series[0]
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.noFill = True
        s.dLbls = DataLabelList()
        s.dLbls.showVal = True
        for a in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
            setattr(s.dLbls, a, False)
        s.dLbls.numFmt = FMT_N
        s.dLbls.position = "outEnd"
        s.dLbls.txPr = _rich(800, INK)
        ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.delete = True
        ch.y_axis.scaling.min = 0
        return ch

    @staticmethod
    def _col(letras):
        from openpyxl.utils import column_index_from_string
        return column_index_from_string(letras)

    def p_ranking(self, cols, fila, alto, completo):
        self.poner(self._barras_h(self.u1, SERIES[0]), cols, fila, alto, completo)
        return ('=IF(cCat1>0,"Ranking por "&Config!$E$9,"Ranking")',
                self.aviso("cCat1=0", "Elige una categoría en Config",
                           f'"Top 10 por "&nomMed&" · líder: "&Motor!${self.u1["lab"]}$3'))

    def p_ranking2(self, cols, fila, alto, completo):
        self.poner(self._barras_h(self.u2, SERIES[6]), cols, fila, alto, completo)
        return ('=IF(cCat2>0,"Ranking por "&Config!$E$10,"Segunda categoría")',
                self.aviso("cCat2=0", "No hay segunda categoría: elígela en Config", '"Top 10 por "&nomMed'))

    def p_dona(self, cols, fila, alto, completo):
        ch = DoughnutChart()
        ch.holeSize = 60
        ch.add_data(Reference(self.mot, min_col=34, min_row=2, max_row=8), titles_from_data=True)     # AH
        ch.set_categories(Reference(self.mot, min_col=33, min_row=3, max_row=8))                     # AG
        ch.style = 2
        ch.title = None
        ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
        ch.plot_area.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
        ch.legend.position = "r"
        ch.legend.txPr = _rich(900, INK_2)
        s = ch.series[0]
        for i in range(6):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = OTROS if i == 5 else SERIES[i]
            pt.graphicalProperties.line.solidFill = CARD
            pt.graphicalProperties.line.width = 25400
            s.dPt.append(pt)
        s.dLbls = DataLabelList()
        s.dLbls.showPercent = True
        for a in ("showVal", "showSerName", "showCatName", "showLegendKey"):
            setattr(s.dLbls, a, False)
        s.dLbls.numFmt = "0%;;;"
        s.dLbls.txPr = _rich(800, INK, True)
        self.poner(ch, cols, fila, alto, completo)
        return ('="Participación por "&IF(dona2,Config!$E$10,Config!$E$9)',
                self.aviso("cCat1=0", "Elige una categoría en Config", '"% de "&nomMed&" · top 5 + Otros"'))

    def p_apiladas(self, cols, fila, alto, completo):
        ch = BarChart()
        ch.type = "col"
        ch.grouping = "stacked"
        ch.overlap = 100
        estilo_chart(ch, leyenda=True)
        ch.add_data(Reference(self.mot, min_col=39, max_col=43, min_row=2, max_row=14), titles_from_data=True)  # AM..AQ
        ch.set_categories(Reference(self.mot, min_col=38, min_row=3, max_row=14))                             # AL
        ch.gapWidth = 40
        for i, s in enumerate(ch.series):
            s.graphicalProperties.solidFill = OTROS if i == 4 else SERIES[i]
            s.graphicalProperties.line.solidFill = CARD
            s.graphicalProperties.line.width = 12700
        ch.y_axis.numFmt = FMT_K
        ch.y_axis.scaling.min = 0
        self.poner(ch, cols, fila, alto, completo)
        return ('="Composición mensual por "&IF(cCat2>0,Config!$E$10,Config!$E$9)',
                self.aviso("cFecha=0", "Necesita columna de fecha (Config)", '"Últimos 12 meses · barras apiladas"'))

    def p_embudo(self, cols, fila, alto, completo):
        ch = BarChart()
        ch.type = "bar"
        ch.grouping = "stacked"
        ch.overlap = 100
        estilo_chart(ch, grid_y=False)
        ch.add_data(Reference(self.mot, min_col=72, max_col=73, min_row=2, max_row=8), titles_from_data=True)  # BT..BU
        ch.set_categories(Reference(self.mot, min_col=71, min_row=3, max_row=8))                              # BS
        ch.gapWidth = 15
        relleno, valor = ch.series
        relleno.graphicalProperties.noFill = True
        relleno.graphicalProperties.line.noFill = True
        valor.graphicalProperties.solidFill = SERIES[2]
        valor.graphicalProperties.line.noFill = True
        valor.dLbls = DataLabelList()
        valor.dLbls.showVal = True
        for a in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
            setattr(valor.dLbls, a, False)
        valor.dLbls.position = "ctr"
        valor.dLbls.numFmt = FMT_N
        valor.dLbls.txPr = _rich(900, "FFFFFF", True)
        ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.delete = True
        self.poner(ch, cols, fila, alto, completo)
        return ('="Embudo"&IF(cEtapa>0," por "&Config!$E$11,"")',
                self.aviso("cEtapa=0", "Elige la columna de Etapa en Config (ej. Visita, Carrito, Compra)",
                           '"Registros por etapa · conversión total: "&ROUND(IFERROR(LOOKUP(2,1/(Motor!$BU$3:$BU$8<>""),'
                           'Motor!$BU$3:$BU$8)/Motor!$BU$3,0)*100,1)&"%"'))

    def p_dispersion(self, cols, fila, alto, completo):
        ch = ScatterChart()
        estilo_chart(ch)
        ch.scatterStyle = "marker"
        xs = Reference(self.mot, min_col=76, min_row=3, max_row=MAX_DISP + 2)   # BX
        ys = Reference(self.mot, min_col=77, min_row=3, max_row=MAX_DISP + 2)   # BY
        s = Series(ys, xs, title="Registros")
        s.marker = Marker(symbol="circle", size=5)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=SERIES[0])
        s.marker.graphicalProperties.line.solidFill = CARD
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
        ch.series.append(s)
        ch.x_axis.numFmt = FMT_K
        ch.y_axis.numFmt = FMT_K
        self.poner(ch, cols, fila, alto, completo)
        return ('="Relación "&nomMed&IF(cMed2>0," vs "&Config!$E$7,"")',
                self.aviso("cMed2=0", "Necesita dos columnas numéricas (Config)",
                           f'"Eje X: "&Config!$E$7&" · eje Y: "&nomMed&" · hasta {MAX_DISP:,} registros"'))

    def p_matriz(self, cols, fila, alto, completo):
        ws = self.dash
        mcols = ["AT", "AU", "AV", "AW", "AX", "AY", "AZ", "BA"]
        # etiqueta en 2 columnas, luego 6 meses y total (se salta la columna separadora)
        destino = [(cols[0], cols[1])] + [(c, c) for c in cols[2:9]]
        r0 = fila + 1
        for (a, b), mc in zip(destino, mcols):
            if b > a:
                for rr in range(r0, r0 + 11):
                    ws.merge_cells(start_row=rr, start_column=a, end_row=rr, end_column=b)
            h = ws.cell(r0, a, f"=Motor!{mc}$2")
            h.font = _font(8, True, INK_2)
            h.alignment = Alignment(horizontal="left" if a == cols[0] else "center")
            h.border = Border(bottom=Side(style="thin", color=INK_3))
            for i in range(10):
                c = ws.cell(r0 + 1 + i, a, f"=Motor!{mc}{i + 3}")
                c.font = _font(9, bold=(mc == "BA"))
                c.number_format = FMT_K
                c.alignment = Alignment(horizontal="left" if a == cols[0] else "center", shrink_to_fit=True)
        v0, v1 = L(destino[1][0]), L(destino[6][0])
        rango = f"{v0}{r0 + 1}:{v1}{r0 + 10}"
        ws.conditional_formatting.add(rango, ColorScaleRule(start_type="min", start_color=SEQ_LO, mid_type="percentile",
                                                            mid_value=50, mid_color=SEQ_MID, end_type="max", end_color=SEQ_HI))
        ws.conditional_formatting.add(rango, FormulaRule(
            formula=[f"AND(ISNUMBER({v0}{r0 + 1}),{v0}{r0 + 1}>=PERCENTILE(${v0}${r0 + 1}:${v1}${r0 + 10},0.7))"],
            font=Font(name=FONT, size=9, bold=True, color="FFFFFF")))
        tot = L(destino[7][0])
        ws.conditional_formatting.add(f"{tot}{r0 + 1}:{tot}{r0 + 10}",
                                      DataBarRule(start_type="num", start_value=0, end_type="max", color="86B6EF"))
        return ('="Matriz: "&IF(cCat1>0,Config!$E$9,"categoría")&" × mes"',
                self.aviso("OR(cFecha=0,cCat1=0)", "Necesita fecha y categoría (Config)",
                           '"Tabla dinámica de "&nomMed&" · últimos 6 meses · color según el valor"'))

    def p_mapa(self, cols, fila, alto, completo):
        ws = self.dash
        r0 = fila
        celdas = []
        n = self.n_deptos
        for i, (nom, ab, _, fr, fc) in enumerate(DEPTOS):
            col = cols[fc - 1]
            ra, rb = r0 + (fr - 1) * 2, r0 + (fr - 1) * 2 + 1
            a = ws.cell(ra, col, f"=N(Motor!$BN${i + 3})")
            a.number_format = f'"{ab}";"{ab}";"{ab}";"{ab}"'
            a.font = _font(8, True, INK)
            a.alignment = Alignment(horizontal="center", vertical="bottom")
            b = ws.cell(rb, col)
            b.value = f'=IF(N(Motor!$BN${i + 3})=0,"",Motor!$BN${i + 3})'
            b.number_format = FMT_K
            b.font = _font(8, color=INK)
            b.alignment = Alignment(horizontal="center", vertical="top", shrink_to_fit=True)
            borde = Side(style="medium", color=CARD)
            a.border = Border(left=borde, right=borde, top=borde)
            b.border = Border(left=borde, right=borde, bottom=borde)
            a.fill = b.fill = PatternFill("solid", fgColor="F0EFEC")
            celdas += [a.coordinate, b.coordinate]
        rango = " ".join(celdas)
        ws.conditional_formatting.add(rango, ColorScaleRule(start_type="min", start_color=SEQ_LO, mid_type="percentile",
                                                            mid_value=50, mid_color=SEQ_MID, end_type="max", end_color=SEQ_HI))
        ws.conditional_formatting.add(rango, FormulaRule(
            formula=[f"AND(ISNUMBER({celdas[0]}),{celdas[0]}>=PERCENTILE(Motor!$BN$3:$BN${n + 2},0.75))"],
            font=Font(name=FONT, size=8, bold=True, color="FFFFFF")))
        return ('="Mapa de Colombia por "&IF(cDepto>0,Config!$E$12,"departamento")',
                self.aviso("cDepto=0", "Elige la columna de Departamento en Config",
                           '"Mosaico por departamento · más oscuro = mayor "&nomMed'))

    def p_top_deptos(self, cols, fila, alto, completo):
        ws = self.dash
        r0 = fila + 1
        for (a, b), h in zip([(cols[0], cols[2]), (cols[3], cols[5]), (cols[6], cols[6])],
                             ["Departamento", "=nomMed", "%"]):
            ws.merge_cells(start_row=r0, start_column=a, end_row=r0, end_column=b)
            c = ws.cell(r0, a, h)
            c.font = _font(8, True, INK_2)
            c.border = Border(bottom=Side(style="thin", color=INK_3))
            c.alignment = Alignment(horizontal="left" if a == cols[0] else "right")
            for k in range(5):
                rr = r0 + 1 + k * 2
                ws.merge_cells(start_row=rr, start_column=a, end_row=rr + 1, end_column=b)
        for k in range(5):
            rr = r0 + 1 + k * 2
            c = ws.cell(rr, cols[0], f"=Motor!$BP${k + 3}")
            c.font = _font(11, True)
            c.alignment = Alignment(vertical="center", indent=1)
            v = ws.cell(rr, cols[3], f"=Motor!$BQ${k + 3}")
            v.number_format = FMT_K
            v.font = _font(11)
            v.alignment = Alignment(vertical="center")
            p = ws.cell(rr, cols[6], f'=IF(Motor!$BQ${k + 3}="","",IFERROR(Motor!$BQ${k + 3}/total,""))')
            p.number_format = "0%"
            p.font = _font(10, color=INK_2)
            p.alignment = Alignment(vertical="center")
        ws.conditional_formatting.add(f"{L(cols[3])}{r0 + 1}:{L(cols[3])}{r0 + 10}",
                                      DataBarRule(start_type="num", start_value=0, end_type="max", color="86B6EF"))
        return '="Top 5 departamentos"', '="Participación sobre el total de "&nomMed'

    def p_calor(self, cols, fila, alto, completo):
        ws = self.dash
        mcols = ["BC"] + [L(56 + d) for d in range(7)]
        destino = [(cols[0], cols[1])] + [(c, c) for c in cols[2:9]]
        r0 = fila
        for (a, b), mc in zip(destino, mcols):
            for i in range(25):
                rr = r0 + i
                if b > a:
                    ws.merge_cells(start_row=rr, start_column=a, end_row=rr, end_column=b)
                c = ws.cell(rr, a, f"=Motor!{mc}{i + 2}")
                c.font = _font(8, bold=(i == 0), color=INK_2 if (i == 0 or mc == "BC") else INK)
                c.number_format = FMT_K
                c.alignment = Alignment(horizontal="right" if mc == "BC" else "center", vertical="center")
                if mc != "BC" and i > 0:
                    s = Side(style="thin", color=CARD)
                    c.border = Border(left=s, right=s, top=s, bottom=s)
            for i in range(25):
                ws.row_dimensions[r0 + i].height = 15
        v0, v1 = L(destino[1][0]), L(destino[7][0])
        rango = f"{v0}{r0 + 1}:{v1}{r0 + 24}"
        ws.conditional_formatting.add(rango, ColorScaleRule(start_type="min", start_color=SEQ_LO, mid_type="percentile",
                                                            mid_value=50, mid_color=SEQ_MID, end_type="max", end_color=SEQ_HI))
        ws.conditional_formatting.add(rango, FormulaRule(
            formula=[f"AND(ISNUMBER({v0}{r0 + 1}),{v0}{r0 + 1}>=PERCENTILE(${v0}${r0 + 1}:${v1}${r0 + 24},0.75))"],
            font=Font(name=FONT, size=8, bold=True, color="FFFFFF")))
        return ('="Mapa de calor: "&IF(conHora,"hora","mes")&" × día de la semana"',
                self.aviso("cFecha=0", "Necesita columna de fecha (Config)",
                           '"Detecta patrones: más oscuro = mayor "&nomMed'))

    # ------------------------------------------------------------------
    def construir(self, salida=SALIDA):
        self.datos_ejemplo()
        self.config()
        self.motor()
        self.dashboard()
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.save(salida)
        return salida


def _horas():
    base = np.array([1, .5, .3, .2, .2, .4, 1, 2, 3.5, 4.5, 5, 5.5, 6, 5.5, 5, 5, 5.5, 6, 7, 7.5, 7, 5, 3.5, 2])
    return base / base.sum()


if __name__ == "__main__":
    print("Plantilla creada:", Plantilla().construir())
