#!/usr/bin/env python3
"""
Generador automático de dashboards en Excel.

Recibe CUALQUIER archivo de datos (.xlsx, .xlsm, .xls, .csv, .txt) y produce un
libro de Excel con:

  * Dashboard   -> tarjetas KPI, tendencia, ranking, participación (dona),
                   composición por período, distribución (histograma),
                   mapa de calor, dispersión y tabla con barras de datos.
  * Datos       -> los datos limpios como Tabla de Excel.
  * Calculos    -> agregaciones con fórmulas (SUMIFS / COUNTIFS) que
                   alimentan las gráficas, así todo se recalcula en Excel.
  * Acerca de   -> qué columnas se detectaron y con qué rol.

El script detecta solo: la hoja con más datos, la fila de encabezados, los
números escritos con cualquier formato (1.234,56 · $1,234.56 · (500)), las
fechas, las columnas de identificadores y las de categorías. Las gráficas que
no aplican a los datos (por ejemplo, tendencia sin columna de fecha) se omiten.

Uso:
    python dashboard_excel.py ventas.xlsx
    python dashboard_excel.py datos.csv -o mi_dashboard.xlsx
    python dashboard_excel.py libro.xlsx --hoja "Hoja2" --medida Ventas --categoria Region
    python dashboard_excel.py            (abre un selector de archivo)
"""
from __future__ import annotations

import argparse
import math
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.table import Table, TableStyleInfo

# --------------------------------------------------------------------------
# Paleta (validada para daltonismo; orden fijo, nunca se recicla)
# --------------------------------------------------------------------------
SERIES = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "6250D6", "E34948"]
OTROS = "B4B2AC"               # gris neutro para "Otros"
AREA_FILL = "CDE2FB"           # paso 100 de la rampa azul
SEQ_LO, SEQ_MID, SEQ_HI = "EEF5FD", "6DA7EC", "104281"  # rampa secuencial azul
GOOD, CRITICAL = "0CA30C", "D03B3B"
INK, INK_2, INK_3 = "0B0B0B", "52514E", "8A8984"
GRID = "E6E5E1"
SURFACE = "F4F4F2"             # fondo del tablero
CARD = "FFFFFF"
FONT = "Arial"

MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
MAX_FILAS_EXCEL = 1_048_575
MAX_PUNTOS_DISPERSION = 5000

HINT_MEDIDA = ["venta", "ingreso", "facturac", "revenue", "sales", "monto", "importe", "total",
               "valor", "amount", "utilidad", "ganancia", "profit", "margen", "costo", "cost",
               "gasto", "cantidad", "unidades", "qty", "quantity", "stock", "existencia", "inventario",
               "saldo", "precio", "price"]
# Medidas que son tasas o valores unitarios: sumarlas no tiene sentido, se prefieren otras.
RE_TASA = re.compile(r"unit|promedio|average|avg|tasa|rate|porcentaje|percent|%|precio|price|^costo$|^cost$|"
                     r"edad|calific|rating|satisf|score|puntaje|latitud|longitud")
HINT_CATEGORIA = ["categoria", "category", "region", "zona", "producto", "product", "linea",
                  "segmento", "segment", "canal", "channel", "ciudad", "city", "pais", "country",
                  "departamento", "estado", "sucursal", "tienda", "store", "vendedor", "cliente",
                  "tipo", "marca", "brand", "area", "sede", "grupo"]
RE_ID = re.compile(r"(^|[^a-z])(id|cod|codigo|code|nit|cedula|documento|telefono|celular|tel|"
                   r"phone|zip|postal|dni|rut|sku|nro|numero|folio|factura|invoice|key|uuid)([^a-z]|$)")
RE_ANIO = re.compile(r"(^|[^a-z])(ano|anio|year|periodo|vigencia)([^a-z]|$)")
TOTALES = {"total", "totales", "gran total", "subtotal", "total general", "grand total"}


def normalizar(texto: str) -> str:
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t.lower()).strip()


# --------------------------------------------------------------------------
# Lectura y limpieza
# --------------------------------------------------------------------------
_RE_LIMPIA_NUM = re.compile(r"[\s $€£¥%]|cop|usd|eur|mxn|ars|clp|pen", re.I)


def parse_numero(v):
    """Convierte casi cualquier representación de número a float (o NaN)."""
    if v is None or isinstance(v, bool):
        return np.nan
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v) if not (isinstance(v, float) and math.isnan(v)) else np.nan
    if isinstance(v, (datetime, date, pd.Timestamp)):
        return np.nan
    s = str(v).strip()
    if not s:
        return np.nan
    negativo = s.startswith("(") and s.endswith(")")
    s = _RE_LIMPIA_NUM.sub("", s.strip("()"))
    if s.endswith("-"):
        negativo, s = True, s[:-1]
    if s.startswith("-"):
        negativo, s = not negativo, s[1:]
    if not s or not re.fullmatch(r"[\d.,]+(e[+-]?\d+)?", s, re.I):
        return np.nan
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        partes = s.split(",")
        s = s.replace(",", ".") if (len(partes) == 2 and len(partes[1]) != 3) else s.replace(",", "")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        x = float(s)
    except ValueError:
        return np.nan
    return -x if negativo else x


def leer_archivo(ruta: Path, hoja: str | None) -> tuple[pd.DataFrame, str]:
    ext = ruta.suffix.lower()
    if ext in (".csv", ".txt", ".tsv"):
        ultimo_error = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                df = pd.read_csv(ruta, sep=None, engine="python", header=None, dtype=str,
                                 encoding=enc, skip_blank_lines=True)
                return df, ruta.stem
            except UnicodeDecodeError as e:
                ultimo_error = e
        raise ultimo_error
    hojas = pd.read_excel(ruta, sheet_name=None, header=None)
    if not hojas:
        raise ValueError("El archivo no contiene hojas.")
    if hoja:
        if hoja not in hojas:
            raise ValueError(f"No existe la hoja '{hoja}'. Hojas disponibles: {', '.join(hojas)}")
        return hojas[hoja], hoja
    nombre = max(hojas, key=lambda h: int(hojas[h].notna().sum().sum()))
    return hojas[nombre], nombre


def es_texto_encabezado(v) -> bool:
    return isinstance(v, str) and v.strip() != "" and math.isnan(parse_numero(v))


def preparar_tabla(raw: pd.DataFrame) -> pd.DataFrame:
    """Encuentra la fila de encabezados, quita filas/columnas vacías y filas de totales."""
    raw = raw.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
    if raw.empty:
        raise ValueError("No se encontraron datos.")
    ancho = raw.shape[1]
    fila_enc = 0
    for i in range(min(30, len(raw))):
        fila = raw.iloc[i]
        no_nulos = int(fila.notna().sum())
        textos = sum(es_texto_encabezado(v) for v in fila)
        if no_nulos >= max(2, 0.6 * ancho) and textos >= 0.6 * no_nulos:
            fila_enc = i
            break
    nombres, vistos = [], {}
    for j, v in enumerate(raw.iloc[fila_enc]):
        n = str(v).strip() if pd.notna(v) and str(v).strip() else f"Columna {j + 1}"
        if isinstance(v, (datetime, pd.Timestamp)):
            n = pd.Timestamp(v).strftime("%Y-%m-%d")
        base, k = n, 2
        while normalizar(n) in vistos:
            n, k = f"{base} ({k})", k + 1
        vistos[normalizar(n)] = True
        nombres.append(n)
    df = raw.iloc[fila_enc + 1:].copy()
    df.columns = nombres
    df = df.dropna(how="all").dropna(axis=1, how="all")
    # Filas de totales (p. ej. "Total general" al final de un reporte)
    es_total = df.apply(lambda r: any(isinstance(v, str) and normalizar(v) in TOTALES for v in r), axis=1)
    df = df[~es_total]
    if len(df) > MAX_FILAS_EXCEL:
        df = df.iloc[:MAX_FILAS_EXCEL]
    return df.reset_index(drop=True)


# --------------------------------------------------------------------------
# Perfilado de columnas
# --------------------------------------------------------------------------
@dataclass
class Perfil:
    df: pd.DataFrame
    tipos: dict = field(default_factory=dict)   # col -> fecha|numero|categoria|id|texto
    medida: str | None = None
    medida2: str | None = None
    fecha: str | None = None
    cat1: str | None = None
    cat2: str | None = None
    cat_dona: str | None = None
    fmt_medida: str = "#,##0"


def _puntaje_hint(nombre: str, hints: list[str]) -> int:
    n = normalizar(nombre)
    for i, h in enumerate(hints):
        if h in n:
            return len(hints) - i
    return 0


def perfilar(df: pd.DataFrame, medida=None, categoria=None, fecha=None) -> Perfil:
    p = Perfil(df=df.copy())
    d = p.df
    n_filas = len(d)
    for c in list(d.columns):
        s = d[c]
        if pd.api.types.is_object_dtype(s) or pd.api.types.is_string_dtype(s):
            s = s.map(lambda v: v.strip() if isinstance(v, str) else v).replace("", np.nan)
        nn = s.dropna()
        if nn.empty:
            d.drop(columns=c, inplace=True)
            continue
        nombre = normalizar(c)
        if pd.api.types.is_datetime64_any_dtype(s):
            d[c], p.tipos[c] = pd.to_datetime(s), "fecha"
            continue
        if pd.api.types.is_bool_dtype(s):
            d[c], p.tipos[c] = s.map(lambda v: "Sí" if v else "No"), "categoria"
            continue
        if nn.map(lambda v: isinstance(v, (datetime, date, pd.Timestamp))).mean() >= 0.8:
            d[c], p.tipos[c] = pd.to_datetime(s, errors="coerce"), "fecha"
            continue
        num = s.map(parse_numero)
        if num.notna().sum() >= 0.9 * len(nn):
            enteros = bool(np.all(np.mod(num.dropna(), 1) == 0))
            if RE_ANIO.search(nombre) and enteros and num.dropna().between(1900, 2100).all():
                d[c], p.tipos[c] = num.map(lambda x: str(int(x)) if pd.notna(x) else np.nan), "categoria"
            elif RE_ID.search(nombre) or (enteros and n_filas > 20 and num.nunique() == num.notna().sum()
                                          and num.dropna().is_monotonic_increasing):
                d[c], p.tipos[c] = s.astype(str).where(s.notna()), "id"
            else:
                d[c], p.tipos[c] = num, "numero"
            continue
        texto = nn.astype(str)
        if texto.str.contains(r"\d", regex=True).mean() >= 0.8:
            fechas = pd.to_datetime(texto, errors="coerce", dayfirst=True, format="mixed")
            if fechas.notna().mean() >= 0.8:
                d[c], p.tipos[c] = pd.to_datetime(s.astype(str).where(s.notna()), errors="coerce",
                                                  dayfirst=True, format="mixed"), "fecha"
                continue
        unicos = texto.nunique()
        if RE_ID.search(nombre) and unicos > 0.8 * len(nn):
            p.tipos[c] = "id"
        elif unicos <= 50 or (unicos <= 500 and unicos <= 0.3 * len(nn)):
            p.tipos[c] = "categoria"
        else:
            p.tipos[c] = "texto"
        d[c] = s.map(lambda v: str(v) if pd.notna(v) else v)

    for c, t in p.tipos.items():
        if t == "categoria":
            d[c] = d[c].fillna("(Sin dato)").astype(str).str.slice(0, 250)

    numeros = [c for c, t in p.tipos.items() if t == "numero"]
    fechas = [c for c, t in p.tipos.items() if t == "fecha"]
    cats = [c for c, t in p.tipos.items() if t == "categoria" and d[c].nunique() >= 2]

    def puntaje_medida(c):
        return (not RE_TASA.search(normalizar(c)), _puntaje_hint(c, HINT_MEDIDA), float(np.nansum(np.abs(d[c]))))

    numeros.sort(key=puntaje_medida, reverse=True)
    p.medida = _elegir(medida, numeros, "medida") if medida else (numeros[0] if numeros else None)
    otras = [c for c in numeros if c != p.medida]
    p.medida2 = otras[0] if otras else None

    fechas.sort(key=lambda c: (d[c].notna().mean(), d[c].nunique()), reverse=True)
    p.fecha = _elegir(fecha, fechas, "fecha") if fecha else (fechas[0] if fechas else None)

    def puntaje_cat(c):
        u = d[c].nunique()
        return (_puntaje_hint(c, HINT_CATEGORIA) > 0, 3 <= u <= 30, _puntaje_hint(c, HINT_CATEGORIA), -abs(u - 8))

    cats.sort(key=puntaje_cat, reverse=True)
    p.cat1 = _elegir(categoria, cats, "categoría") if categoria else (cats[0] if cats else None)
    resto = [c for c in cats if c != p.cat1]
    p.cat2 = resto[0] if resto else None
    pocas = sorted([c for c in cats if 2 <= d[c].nunique() <= 6], key=lambda c: (c == p.cat1, d[c].nunique()))
    p.cat_dona = pocas[0] if pocas else p.cat1

    if p.medida is not None:
        vals = d[p.medida].dropna()
        p.fmt_medida = "#,##0" if (np.all(np.mod(vals, 1) == 0) or vals.abs().mean() >= 1000) else "#,##0.00"
    return p


def _elegir(pedido: str, opciones: list[str], que: str) -> str:
    for c in opciones:
        if normalizar(c) == normalizar(pedido):
            return c
    raise ValueError(f"La {que} '{pedido}' no se encontró o no es del tipo correcto. Opciones: {opciones}")


# --------------------------------------------------------------------------
# Utilidades de estilo
# --------------------------------------------------------------------------
def _font(size=10, bold=False, color=INK):
    return Font(name=FONT, size=size, bold=bold, color=color)


def _rich(size=900, color=INK_2, bold=False):
    cp = CharacterProperties(latin=DFont(typeface=FONT), sz=size, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _linea(color=None, ancho_pt=0.75):
    if color is None:
        return GraphicalProperties(ln=LineProperties(noFill=True))
    return GraphicalProperties(ln=LineProperties(solidFill=color, w=int(ancho_pt * 12700)))


def estilo_chart(ch, leyenda=False, grid_y=True):
    ch.style = 2
    ch.roundedCorners = False
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.graphical_properties.solidFill = CARD
    ch.title = None
    if leyenda:
        ch.legend.position = "t"
        ch.legend.txPr = _rich(900, INK_2)
    else:
        ch.legend = None
    for eje in (getattr(ch, "x_axis", None), getattr(ch, "y_axis", None)):
        if eje is None:
            continue
        eje.delete = False
        eje.txPr = _rich(800, INK_2)
        eje.spPr = _linea(GRID)
        eje.majorTickMark = "none"
        eje.minorTickMark = "none"
        eje.majorGridlines = None
    if grid_y and hasattr(ch, "y_axis"):
        from openpyxl.chart.axis import ChartLines
        ch.y_axis.majorGridlines = ChartLines(spPr=_linea(GRID, 0.5))
        ch.y_axis.spPr = _linea(None)


def _rango(ws, col, f0, f1):
    return f"{quote_sheetname(ws.title)}!${get_column_letter(col)}${f0}:${get_column_letter(col)}${f1}"


def _criterio_texto(celda: str, valor: str) -> str:
    """Criterio de SUMIFS que coincide exactamente con el texto (escapa comodines)."""
    expr = celda
    if any(ch in valor for ch in "~*?"):
        expr = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({celda},"~","~~"),"*","~*"),"?","~?")'
    if valor[:1] in "=<>" or expr != celda:
        expr = f'"="&{expr}'
    return expr


# --------------------------------------------------------------------------
# Constructor del libro
# --------------------------------------------------------------------------
class Constructor:
    def __init__(self, perfil: Perfil, titulo: str, origen: str):
        self.p = perfil
        self.titulo = titulo
        self.origen = origen
        self.wb = Workbook()
        self.dash = self.wb.active
        self.dash.title = "Dashboard"
        self.datos = self.wb.create_sheet("Datos")
        self.calc = self.wb.create_sheet("Calculos")
        self.info = self.wb.create_sheet("Acerca de")
        self.fila_calc = 1
        self.col_ref: dict[str, str] = {}   # columna -> rango absoluto en Datos
        self.paneles: list[dict] = []
        self.kpis: list[dict] = []

    # ---------------- Datos ----------------
    def escribir_datos(self):
        ws, d = self.datos, self.p.df
        ws.append(list(d.columns))
        fmts = {}
        for c in d.columns:
            t = self.p.tipos[c]
            fmts[c] = ("dd/mm/yyyy" if t == "fecha" else
                       (self.p.fmt_medida if c == self.p.medida else "#,##0.##") if t == "numero" else None)
        conv = []
        for c in d.columns:
            t = self.p.tipos[c]
            if t == "fecha":
                conv.append(lambda v: v.to_pydatetime() if pd.notna(v) else None)
            elif t == "numero":
                conv.append(lambda v: float(v) if pd.notna(v) else None)
            else:
                conv.append(lambda v: v if (v is not None and not (isinstance(v, float) and math.isnan(v))) else None)
        for fila in d.itertuples(index=False, name=None):
            ws.append([f(v) for f, v in zip(conv, fila)])
        n = len(d) + 1
        for j, c in enumerate(d.columns, start=1):
            letra = get_column_letter(j)
            self.col_ref[c] = f"Datos!${letra}$2:${letra}${n}"
            if fmts[c]:
                for (celda,) in ws.iter_rows(min_row=2, max_row=n, min_col=j, max_col=j):
                    celda.number_format = fmts[c]
            largo = max([len(str(c))] + [len(str(v)) for v in d[c].head(200).tolist()])
            ws.column_dimensions[letra].width = min(max(10, largo + 2), 45)
        tabla = Table(displayName="TablaDatos", ref=f"A1:{get_column_letter(len(d.columns))}{n}")
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tabla)
        ws.freeze_panes = "A2"
        self.n = n

    # ---------------- helpers de Calculos ----------------
    def agg(self, criterios: list[tuple[str, str]]) -> str:
        partes = ",".join(f"{r},{c}" for r, c in criterios)
        if self.p.medida:
            return f"SUMIFS({self.col_ref[self.p.medida]},{partes})" if partes else f"SUM({self.col_ref[self.p.medida]})"
        return f"COUNTIFS({partes})" if partes else f"COUNTA({next(iter(self.col_ref.values()))})"

    @property
    def nombre_medida(self):
        return self.p.medida or "Registros"

    @property
    def fmt(self):
        return self.p.fmt_medida if self.p.medida else "#,##0"

    def bloque(self, titulo: str, encabezados: list[str]) -> int:
        ws = self.calc
        ws.cell(self.fila_calc, 1, titulo).font = _font(11, True)
        for j, h in enumerate(encabezados, start=1):
            c = ws.cell(self.fila_calc + 1, j, h)
            c.font = _font(9, True, "FFFFFF")
            c.fill = PatternFill("solid", fgColor="2A78D6")
        inicio = self.fila_calc + 2
        return inicio

    def cerrar_bloque(self, ultima_fila: int):
        self.fila_calc = ultima_fila + 3

    # ---------------- análisis ----------------
    def analisis_ranking(self, cat: str, top: int = 10) -> dict:
        d, ws = self.p.df, self.calc
        if self.p.medida:
            orden = d.groupby(cat)[self.p.medida].sum().sort_values(ascending=False)
        else:
            orden = d[cat].value_counts()
        etiquetas = [str(x) for x in orden.index[:top]]
        f0 = self.bloque(f"Ranking de {cat} por {self.nombre_medida}", [cat, self.nombre_medida, "% del total"])
        total_ref = f"SUM({self.col_ref[self.p.medida]})" if self.p.medida else f"ROWS({self.col_ref[cat]})"
        for i, et in enumerate(etiquetas):
            r = f0 + i
            ws.cell(r, 1, et)
            ws.cell(r, 2, "=" + self.agg([(self.col_ref[cat], _criterio_texto(f"$A{r}", et))])).number_format = self.fmt
            ws.cell(r, 3, f"=IFERROR(B{r}/{total_ref},0)").number_format = "0.0%"
        f1 = f0 + len(etiquetas) - 1
        self.cerrar_bloque(f1)
        return {"f0": f0, "f1": f1, "n_total": len(orden), "cat": cat}

    def analisis_dona(self, cat: str, maximo: int = 5) -> dict:
        d, ws = self.p.df, self.calc
        orden = (d.groupby(cat)[self.p.medida].sum() if self.p.medida else d[cat].value_counts())
        orden = orden.sort_values(ascending=False)
        etiquetas = [str(x) for x in orden.index[:maximo]]
        f0 = self.bloque(f"Participación por {cat}", [cat, self.nombre_medida])
        for i, et in enumerate(etiquetas):
            r = f0 + i
            ws.cell(r, 1, et)
            ws.cell(r, 2, "=" + self.agg([(self.col_ref[cat], _criterio_texto(f"$A{r}", et))])).number_format = self.fmt
        f1 = f0 + len(etiquetas) - 1
        if len(orden) > maximo:
            f1 += 1
            ws.cell(f1, 1, "Otros")
            ws.cell(f1, 2, f"={self.agg([])}-SUM(B{f0}:B{f1 - 1})").number_format = self.fmt
        self.cerrar_bloque(f1)
        return {"f0": f0, "f1": f1, "otros": len(orden) > maximo, "cat": cat}

    def periodos(self) -> tuple[str, list[tuple[str, datetime, datetime]]]:
        s = self.p.df[self.p.fecha].dropna()
        ini, fin = s.min().normalize(), s.max().normalize()
        dias = (fin - ini).days
        out = []
        if dias <= 62:
            gran, cur = "Día", ini
            while cur <= fin:
                out.append((f"{cur.day:02d} {MESES[cur.month - 1]}", cur, cur + timedelta(days=1)))
                cur += timedelta(days=1)
        elif dias <= 182:
            gran, cur = "Semana", ini - timedelta(days=ini.weekday())
            while cur <= fin:
                out.append((f"{cur.day:02d} {MESES[cur.month - 1]}", cur, cur + timedelta(days=7)))
                cur += timedelta(days=7)
        elif dias <= 4 * 366:
            gran, cur = "Mes", pd.Timestamp(ini.year, ini.month, 1)
            while cur <= fin:
                sig = cur + pd.offsets.MonthBegin(1)
                out.append((f"{MESES[cur.month - 1]} {str(cur.year)[2:]}", cur, sig))
                cur = sig
        elif dias <= 12 * 366:
            gran, cur = "Trimestre", pd.Timestamp(ini.year, 3 * ((ini.month - 1) // 3) + 1, 1)
            while cur <= fin:
                sig = cur + pd.offsets.MonthBegin(3)
                out.append((f"T{(cur.month - 1) // 3 + 1} {cur.year}", cur, sig))
                cur = sig
        else:
            gran, cur = "Año", pd.Timestamp(ini.year, 1, 1)
            while cur <= fin:
                sig = pd.Timestamp(cur.year + 1, 1, 1)
                out.append((str(cur.year), cur, sig))
                cur = sig
        return gran, [(e, a.to_pydatetime(), b.to_pydatetime()) for e, a, b in out]

    def analisis_tendencia(self) -> dict:
        ws = self.calc
        gran, per = self.periodos()
        f0 = self.bloque(f"{self.nombre_medida} por {gran.lower()}", ["Período", "Desde", "Hasta (excl.)", self.nombre_medida])
        rf = self.col_ref[self.p.fecha]
        for i, (et, a, b) in enumerate(per):
            r = f0 + i
            ws.cell(r, 1, et)
            ws.cell(r, 2, a).number_format = "dd/mm/yyyy"
            ws.cell(r, 3, b).number_format = "dd/mm/yyyy"
            ws.cell(r, 4, "=" + self.agg([(rf, f'">="&$B{r}'), (rf, f'"<"&$C{r}')])).number_format = self.fmt
        f1 = f0 + len(per) - 1
        self.cerrar_bloque(f1)
        # El último período suele estar incompleto (p. ej. mes en curso): el KPI compara completos.
        ultimo_dato = self.p.df[self.p.fecha].max().to_pydatetime()
        incompleto = ultimo_dato < per[-1][2] - timedelta(days=1)
        f_kpi = f1 - 1 if (incompleto and len(per) >= 3) else f1
        return {"f0": f0, "f1": f1, "gran": gran, "f_kpi": f_kpi, "completo": f_kpi != f1 or not incompleto}

    def analisis_composicion(self, cat: str, tend: dict, maximo: int = 4) -> dict:
        d, ws = self.p.df, self.calc
        orden = (d.groupby(cat)[self.p.medida].sum() if self.p.medida else d[cat].value_counts()).sort_values(ascending=False)
        etiquetas = [str(x) for x in orden.index[:maximo]]
        otros = len(orden) > maximo
        enc = ["Período"] + etiquetas + (["Otros"] if otros else [])
        f0 = self.bloque(f"{self.nombre_medida} por {tend['gran'].lower()} y {cat}", enc)
        rf, rc = self.col_ref[self.p.fecha], self.col_ref[cat]
        enc_fila = f0 - 1
        for i in range(tend["f1"] - tend["f0"] + 1):
            r, rt = f0 + i, tend["f0"] + i
            ws.cell(r, 1, f"=A{rt}")
            for j, et in enumerate(etiquetas, start=2):
                col = get_column_letter(j)
                crit = [(rf, f'">="&$B{rt}'), (rf, f'"<"&$C{rt}'), (rc, _criterio_texto(f"{col}${enc_fila}", et))]
                ws.cell(r, j, "=" + self.agg(crit)).number_format = self.fmt
            if otros:
                j = len(etiquetas) + 2
                ws.cell(r, j, f"=D{rt}-SUM(B{r}:{get_column_letter(j - 1)}{r})").number_format = self.fmt
        f1 = f0 + (tend["f1"] - tend["f0"])
        self.cerrar_bloque(f1)
        return {"f0": f0, "f1": f1, "n_series": len(enc) - 1, "otros": otros, "cat": cat, "enc_fila": enc_fila}

    def analisis_histograma(self) -> dict | None:
        vals = self.p.df[self.p.medida].dropna()
        if vals.nunique() < 5:
            return None
        lo, hi = float(np.percentile(vals, 0)), float(np.percentile(vals, 100))
        if hi <= lo:
            return None
        bruto = (hi - lo) / 10
        mag = 10 ** math.floor(math.log10(bruto))
        paso = min((m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= bruto), default=10 * mag)
        inicio = math.floor(lo / paso) * paso
        bordes = [inicio]
        while bordes[-1] < hi:
            bordes.append(bordes[-1] + paso)
        ws = self.calc
        f0 = self.bloque(f"Distribución de {self.p.medida}", ["Rango", "Desde", "Hasta", "Frecuencia"])
        rm = self.col_ref[self.p.medida]
        dec = 0 if paso >= 1 else 2
        for i in range(len(bordes) - 1):
            r = f0 + i
            a, b = bordes[i], bordes[i + 1]
            ws.cell(r, 1, f"{a:,.{dec}f} – {b:,.{dec}f}")
            ws.cell(r, 2, a)
            ws.cell(r, 3, b)
            op = '"<="' if i == len(bordes) - 2 else '"<"'
            ws.cell(r, 4, f'=COUNTIFS({rm},">="&$B{r},{rm},{op}&$C{r})')
        f1 = f0 + len(bordes) - 2
        self.cerrar_bloque(f1)
        return {"f0": f0, "f1": f1}

    def analisis_calor_cats(self, ca: str, cb: str) -> dict:
        d, ws = self.p.df, self.calc
        def top(c, k):
            o = d.groupby(c)[self.p.medida].sum() if self.p.medida else d[c].value_counts()
            return [str(x) for x in o.sort_values(ascending=False).index[:k]]
        filas, cols = top(ca, 10), top(cb, 6)
        f0 = self.bloque(f"{self.nombre_medida}: {ca} × {cb}", [f"{ca} \\ {cb}"] + cols)
        enc = f0 - 1
        for i, fa in enumerate(filas):
            r = f0 + i
            ws.cell(r, 1, fa)
            for j, cbv in enumerate(cols, start=2):
                col = get_column_letter(j)
                crit = [(self.col_ref[ca], _criterio_texto(f"$A{r}", fa)), (self.col_ref[cb], _criterio_texto(f"{col}${enc}", cbv))]
                ws.cell(r, j, "=" + self.agg(crit)).number_format = self.fmt
        f1 = f0 + len(filas) - 1
        self.cerrar_bloque(f1)
        return {"f0": f0, "f1": f1, "enc": enc, "ncols": len(cols) + 1,
                "titulo": f"{ca} × {cb}", "sub": f"{self.nombre_medida} — intensidad de color = mayor valor"}

    def analisis_calor_fechas(self) -> dict | None:
        s = self.p.df[self.p.fecha].dropna()
        anios = sorted(s.dt.year.unique())
        if (s.max() - s.min()).days < 180:
            return None
        ws = self.calc
        f0 = self.bloque(f"{self.nombre_medida} por año y mes", ["Año"] + MESES)
        rf = self.col_ref[self.p.fecha]
        for i, a in enumerate(anios):
            r = f0 + i
            ws.cell(r, 1, int(a))
            for m in range(1, 13):
                ini, fin = f"DATE($A{r},{m},1)", f"DATE($A{r},{m + 1},1)"
                ws.cell(r, m + 1, "=" + self.agg([(rf, f'">="&{ini}'), (rf, f'"<"&{fin}')])).number_format = self.fmt
        f1 = f0 + len(anios) - 1
        self.cerrar_bloque(f1)
        return {"f0": f0, "f1": f1, "enc": f0 - 1, "ncols": 13,
                "titulo": "Estacionalidad: año × mes", "sub": f"{self.nombre_medida} por mes — más oscuro = mayor valor"}

    # ---------------- gráficas ----------------
    def _base_cero(self, ch):
        """Barras y áreas siempre arrancan en cero (si no hay negativos)."""
        if not self.p.medida or self.p.df[self.p.medida].min() >= 0:
            ch.y_axis.scaling.min = 0

    def chart_tendencia(self, t):
        ch = AreaChart()
        estilo_chart(ch)
        datos = Reference(self.calc, min_col=4, min_row=t["f0"] - 1, max_row=t["f1"])
        ch.add_data(datos, titles_from_data=True)
        ch.set_categories(Reference(self.calc, min_col=1, min_row=t["f0"], max_row=t["f1"]))
        s = ch.series[0]
        s.graphicalProperties.solidFill = AREA_FILL
        s.graphicalProperties.line.solidFill = SERIES[0]
        s.graphicalProperties.line.width = 28575
        ch.y_axis.numFmt = self.fmt
        ch.y_axis.number_format = self.fmt
        self._base_cero(ch)
        ch.x_axis.tickLblSkip = max(1, (t["f1"] - t["f0"] + 1) // 14)
        return ch

    def chart_ranking(self, r):
        ch = BarChart()
        ch.type = "bar"
        estilo_chart(ch, grid_y=False)
        ch.add_data(Reference(self.calc, min_col=2, min_row=r["f0"] - 1, max_row=r["f1"]), titles_from_data=True)
        ch.set_categories(Reference(self.calc, min_col=1, min_row=r["f0"], max_row=r["f1"]))
        ch.gapWidth = 45
        s = ch.series[0]
        s.graphicalProperties.solidFill = SERIES[0]
        s.graphicalProperties.line.noFill = True
        s.dLbls = DataLabelList()
        s.dLbls.showVal = True
        for k in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
            setattr(s.dLbls, k, False)
        s.dLbls.numFmt = self.fmt
        s.dLbls.position = "outEnd"
        s.dLbls.txPr = _rich(800, INK)
        ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.delete = True
        return ch

    def chart_dona(self, dn):
        ch = DoughnutChart()
        ch.holeSize = 62
        ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
        ch.plot_area.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
        ch.add_data(Reference(self.calc, min_col=2, min_row=dn["f0"] - 1, max_row=dn["f1"]), titles_from_data=True)
        ch.set_categories(Reference(self.calc, min_col=1, min_row=dn["f0"], max_row=dn["f1"]))
        ch.title = None
        ch.style = 2
        ch.legend.position = "r"
        ch.legend.txPr = _rich(900, INK_2)
        s = ch.series[0]
        n = dn["f1"] - dn["f0"] + 1
        for i in range(n):
            pt = DataPoint(idx=i)
            color = OTROS if (dn["otros"] and i == n - 1) else SERIES[i]
            pt.graphicalProperties.solidFill = color
            pt.graphicalProperties.line.solidFill = CARD
            pt.graphicalProperties.line.width = 25400
            s.dPt.append(pt)
        s.dLbls = DataLabelList()
        s.dLbls.showPercent = True
        for k in ("showVal", "showSerName", "showCatName", "showLegendKey"):
            setattr(s.dLbls, k, False)
        s.dLbls.numFmt = "0%"
        s.dLbls.txPr = _rich(800, INK, True)
        return ch

    def chart_composicion(self, c):
        ch = BarChart()
        ch.type = "col"
        ch.grouping = "stacked"
        ch.overlap = 100
        estilo_chart(ch, leyenda=True)
        ch.add_data(Reference(self.calc, min_col=2, max_col=1 + c["n_series"], min_row=c["enc_fila"], max_row=c["f1"]),
                    titles_from_data=True)
        ch.set_categories(Reference(self.calc, min_col=1, min_row=c["f0"], max_row=c["f1"]))
        ch.gapWidth = 40
        for i, s in enumerate(ch.series):
            ultimo = c["otros"] and i == len(ch.series) - 1
            s.graphicalProperties.solidFill = OTROS if ultimo else SERIES[i]
            s.graphicalProperties.line.solidFill = CARD
            s.graphicalProperties.line.width = 12700
        ch.y_axis.numFmt = self.fmt
        ch.y_axis.number_format = self.fmt
        self._base_cero(ch)
        ch.x_axis.tickLblSkip = max(1, (c["f1"] - c["f0"] + 1) // 14)
        return ch

    def chart_histograma(self, h):
        ch = BarChart()
        ch.type = "col"
        estilo_chart(ch)
        ch.add_data(Reference(self.calc, min_col=4, min_row=h["f0"] - 1, max_row=h["f1"]), titles_from_data=True)
        ch.set_categories(Reference(self.calc, min_col=1, min_row=h["f0"], max_row=h["f1"]))
        ch.gapWidth = 8
        s = ch.series[0]
        s.graphicalProperties.solidFill = SERIES[2]
        s.graphicalProperties.line.noFill = True
        ch.y_axis.numFmt = "#,##0"
        ch.y_axis.number_format = "#,##0"
        ch.y_axis.scaling.min = 0
        return ch

    def chart_dispersion(self):
        ch = ScatterChart()
        estilo_chart(ch)
        ch.scatterStyle = "marker"
        x, y = self.p.medida2, self.p.medida
        ix = list(self.p.df.columns).index(x) + 1
        iy = list(self.p.df.columns).index(y) + 1
        fin = min(self.n, MAX_PUNTOS_DISPERSION + 1)
        xs = Reference(self.datos, min_col=ix, min_row=2, max_row=fin)
        ys = Reference(self.datos, min_col=iy, min_row=2, max_row=fin)
        s = Series(ys, xs, title=f"{y} vs {x}")
        s.marker = Marker(symbol="circle", size=6)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=SERIES[0])
        s.marker.graphicalProperties.line.solidFill = CARD
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
        ch.series.append(s)
        ch.x_axis.title = x
        ch.y_axis.title = y
        ch.x_axis.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(latin=DFont(typeface=FONT), sz=900, b=False, solidFill=INK_2))
        ch.y_axis.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(latin=DFont(typeface=FONT), sz=900, b=False, solidFill=INK_2))
        ch.x_axis.numFmt = "#,##0"
        ch.y_axis.numFmt = "#,##0"
        return ch

    # ---------------- construir ----------------
    def construir(self):
        p = self.p
        self.escribir_datos()
        tend = self.analisis_tendencia() if p.fecha else None
        rank = self.analisis_ranking(p.cat1) if p.cat1 else None
        dona = self.analisis_dona(p.cat_dona) if p.cat_dona else None
        comp = None
        if tend and p.cat_dona and tend["f1"] - tend["f0"] >= 1:
            comp = self.analisis_composicion(p.cat_dona, tend)
        hist = self.analisis_histograma() if p.medida else None
        calor = None
        if p.cat1 and p.cat2:
            calor = self.analisis_calor_cats(p.cat1, p.cat2)
        elif p.fecha:
            calor = self.analisis_calor_fechas()
        rank2 = self.analisis_ranking(p.cat2) if p.cat2 else None

        # KPIs
        if p.medida:
            rm = self.col_ref[p.medida]
            self.kpis.append({"t": f"{p.medida} total", "f": f"=SUM({rm})", "fmt": "k"})
            self.kpis.append({"t": f"{p.medida} promedio", "f": f"=IFERROR(AVERAGE({rm}),0)", "fmt": "k"})
        self.kpis.append({"t": "Registros", "f": f"=ROWS({next(iter(self.col_ref.values()))})", "fmt": "#,##0"})
        if tend and tend["f_kpi"] > tend["f0"]:
            fk = tend["f_kpi"]
            u, a = f"Calculos!$D${fk}", f"Calculos!$D${fk - 1}"
            completo = " completo" if tend["completo"] and fk != tend["f1"] else ""
            self.kpis.append({"t": f"Último {tend['gran'].lower()}{completo} vs anterior", "f": f"=IFERROR({u}/{a}-1,0)",
                              "fmt": '▲ 0.0%;▼ 0.0%;0.0%', "delta": True,
                              "sub": f'=Calculos!$A${fk}&": "&TEXT({u},"#,##0")'})
        elif rank:
            self.kpis.append({"t": f"{p.cat1} líder", "f": f"=Calculos!$A${rank['f0']}", "fmt": "@",
                              "sub": f'=TEXT(Calculos!$C${rank["f0"]},"0%")&" del total"'})
        elif p.medida:
            self.kpis.append({"t": f"{p.medida} máximo", "f": f"=MAX({self.col_ref[p.medida]})", "fmt": "k"})

        # Paneles (orden de lectura: tendencia -> quién -> cómo se reparte -> detalle)
        if tend:
            self.paneles.append({"titulo": f"Tendencia de {self.nombre_medida}", "ancho": 2,
                                 "sub": f"Total por {tend['gran'].lower()}", "chart": self.chart_tendencia(tend)})
        if rank:
            sub = f'="Top {min(10, rank["n_total"])} de {rank["n_total"]} · líder: "&Calculos!$A${rank["f0"]}&" ("&TEXT(Calculos!$C${rank["f0"]},"0%")&")"'
            self.paneles.append({"titulo": f"Ranking por {p.cat1}", "sub": sub, "chart": self.chart_ranking(rank)})
        if dona:
            self.paneles.append({"titulo": f"Participación por {p.cat_dona}",
                                 "sub": f"% de {self.nombre_medida}", "chart": self.chart_dona(dona)})
        if comp:
            self.paneles.append({"titulo": f"Composición por {p.cat_dona}",
                                 "sub": f"{self.nombre_medida} por {tend['gran'].lower()}, apilado", "chart": self.chart_composicion(comp)})
        if hist:
            self.paneles.append({"titulo": f"Distribución de {p.medida}",
                                 "sub": "Número de registros por rango de valor", "chart": self.chart_histograma(hist)})
        if calor:
            self.paneles.append({"titulo": calor["titulo"], "sub": calor["sub"], "calor": calor})
        if p.medida and p.medida2:
            nota = f" (primeros {MAX_PUNTOS_DISPERSION:,} registros)" if self.n - 1 > MAX_PUNTOS_DISPERSION else ""
            self.paneles.append({"titulo": f"Relación {p.medida} vs {p.medida2}",
                                 "sub": f"Cada punto es un registro{nota}", "chart": self.chart_dispersion()})
        if rank2:
            self.paneles.append({"titulo": f"Ranking por {p.cat2}", "sub": f"Top {min(10, rank2['n_total'])} por {self.nombre_medida}",
                                 "chart": self.chart_ranking(rank2)})
        if rank:
            self.paneles.append({"titulo": f"Detalle por {p.cat1}", "sub": "Valor y participación con barras de datos", "tabla": rank})

        self.dibujar_dashboard()
        self.escribir_info()
        self.formatear_calc()
        self.wb.calculation.fullCalcOnLoad = True
        return self.wb

    # ---------------- maquetación del Dashboard ----------------
    def dibujar_dashboard(self):
        ws = self.dash
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90
        # Rejilla: A | B..G panel izq | H separador | I..N panel der | O
        anchos = {"A": 2.5, "H": 2.5, "O": 2.5}
        for j in range(1, 16):
            letra = get_column_letter(j)
            ws.column_dimensions[letra].width = anchos.get(letra, 13)
        fondo = PatternFill("solid", fgColor=SURFACE)
        alto_panel = 17
        filas_paneles = 0
        col_actual = 0
        for pnl in self.paneles:
            if pnl.get("ancho") == 2 or col_actual == 0:
                filas_paneles += 1
                col_actual = 0 if pnl.get("ancho") == 2 else 1
            else:
                col_actual = 0
        ultima_fila = 10 + filas_paneles * (alto_panel + 4) + 4
        for fila in ws.iter_rows(min_row=1, max_row=ultima_fila, min_col=1, max_col=15):
            for c in fila:
                c.fill = fondo

        # Encabezado
        ws.row_dimensions[2].height = 30
        ws["B2"] = self.titulo
        ws["B2"].font = _font(20, True)
        ws["B3"] = (f"Fuente: {self.origen} · {self.n - 1:,} registros · "
                    f"generado {datetime.now():%d/%m/%Y %H:%M}")
        ws["B3"].font = _font(9, color=INK_2)

        # Tarjetas KPI: 4 tarjetas de 3 columnas cada una en B..N (saltando H)
        ranuras = [("B", "D"), ("E", "G"), ("I", "K"), ("L", "N")]
        blanco = PatternFill("solid", fgColor=CARD)
        borde = Side(style="thin", color=GRID)
        for (c0, c1), k in zip(ranuras, self.kpis[:4]):
            for r in range(5, 9):
                for col in range(ws[c0 + "1"].column, ws[c1 + "1"].column + 1):
                    celda = ws.cell(r, col)
                    celda.fill = blanco
                    celda.border = Border(top=borde if r == 5 else None, bottom=borde if r == 8 else None,
                                          left=borde if col == ws[c0 + "1"].column else None,
                                          right=borde if col == ws[c1 + "1"].column else None)
            ws.merge_cells(f"{c0}5:{c1}5")
            ws.merge_cells(f"{c0}6:{c1}7")
            ws.merge_cells(f"{c0}8:{c1}8")
            ws[f"{c0}5"] = k["t"].upper()
            ws[f"{c0}5"].font = _font(8, True, INK_2)
            ws[f"{c0}5"].alignment = Alignment(indent=1, vertical="bottom")
            v = ws[f"{c0}6"]
            v.value = k["f"]
            v.font = _font(22, True)
            v.alignment = Alignment(horizontal="left", indent=1, vertical="center", shrink_to_fit=True)
            v.number_format = ('[>=1000000]#,##0.0,,"M";[>=1000]#,##0.0,"K";#,##0.##'
                               if k["fmt"] == "k" else k["fmt"])
            if k.get("delta"):
                ws.conditional_formatting.add(f"{c0}6", FormulaRule(formula=[f"{c0}6>0"], font=Font(color=GOOD, bold=True, size=22, name=FONT)))
                ws.conditional_formatting.add(f"{c0}6", FormulaRule(formula=[f"{c0}6<0"], font=Font(color=CRITICAL, bold=True, size=22, name=FONT)))
            if k.get("sub"):
                ws[f"{c0}8"] = k["sub"]
                ws[f"{c0}8"].font = _font(9, color=INK_2)
                ws[f"{c0}8"].alignment = Alignment(indent=1, vertical="top")
        for r, h in ((5, 18), (6, 22), (7, 22), (8, 18)):
            ws.row_dimensions[r].height = h

        # Paneles
        fila = 11
        col_actual = 0
        ancho_col_px = 13 * 7 + 5
        sep_px = int(2.5 * 7 + 5)
        cm = 0.026458
        for pnl in self.paneles:
            completo = pnl.get("ancho") == 2
            if completo and col_actual == 1:
                fila += alto_panel + 4
                col_actual = 0
            c0, c1 = ("B", "N") if completo else (("B", "G") if col_actual == 0 else ("I", "N"))
            # tarjeta blanca
            n0, n1 = ws[c0 + "1"].column, ws[c1 + "1"].column
            for r in range(fila, fila + alto_panel + 3):
                for col in range(n0, n1 + 1):
                    celda = ws.cell(r, col)
                    celda.fill = blanco
                    celda.border = Border(top=borde if r == fila else None, bottom=borde if r == fila + alto_panel + 2 else None,
                                          left=borde if col == n0 else None, right=borde if col == n1 else None)
            ws.cell(fila, n0, pnl["titulo"]).font = _font(12, True)
            ws.cell(fila, n0).alignment = Alignment(indent=1, vertical="bottom")
            ws.cell(fila + 1, n0, pnl["sub"]).font = _font(9, color=INK_2)
            ws.cell(fila + 1, n0).alignment = Alignment(indent=1)
            ws.row_dimensions[fila].height = 22
            if "chart" in pnl:
                ch = pnl["chart"]
                ancho_px = (12 * ancho_col_px + sep_px) if completo else 6 * ancho_col_px
                ch.width = (ancho_px - 16) * cm
                ch.height = (alto_panel * 20) * cm
                ws.add_chart(ch, f"{c0}{fila + 2}")
            elif "calor" in pnl:
                self.dibujar_calor(pnl["calor"], fila + 3, n0, n1)
            elif "tabla" in pnl:
                self.dibujar_tabla(pnl["tabla"], fila + 3, n0)
            if completo:
                fila += alto_panel + 4
            else:
                col_actual = 1 - col_actual
                if col_actual == 0:
                    fila += alto_panel + 4
        ws.print_area = f"A1:O{ultima_fila}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    def dibujar_calor(self, c, fila, n0, n1):
        """Mapa de calor con celdas: cada columna de Calculos se reparte en el panel."""
        ws = self.dash
        ncols = c["ncols"]
        espacio = n1 - n0 + 1
        # Si hay más columnas que celdas en el panel (p. ej. 12 meses en 6 celdas), se usan dos bandas.
        bandas = [list(range(1, ncols))] if ncols - 1 <= espacio - 1 else [list(range(1, 1 + math.ceil((ncols - 1) / 2))),
                                                                       list(range(1 + math.ceil((ncols - 1) / 2), ncols))]
        filas = c["f1"] - c["f0"] + 1
        r = fila
        rango_valores = []
        for banda in bandas:
            ws.cell(r, n0).value = f"=Calculos!$A${c['enc']}"
            ws.cell(r, n0).font = _font(8, True, INK_2)
            for k, jc in enumerate(banda):
                cel = ws.cell(r, n0 + 1 + k, f"=Calculos!{get_column_letter(jc + 1)}${c['enc']}")
                cel.font = _font(8, True, INK_2)
                cel.alignment = Alignment(horizontal="center")
            for i in range(filas):
                rr = r + 1 + i
                cel = ws.cell(rr, n0, f"=Calculos!$A${c['f0'] + i}")
                cel.font = _font(9, color=INK)
                cel.alignment = Alignment(shrink_to_fit=True, horizontal="left")
                for k, jc in enumerate(banda):
                    v = ws.cell(rr, n0 + 1 + k, f"=Calculos!{get_column_letter(jc + 1)}{c['f0'] + i}")
                    v.number_format = '[>=1000000]#,##0.0,,"M";[>=1000]#,##0.0,"K";#,##0'
                    v.font = _font(9)
                    v.alignment = Alignment(horizontal="center", vertical="center")
                    v.border = Border(left=Side(style="thin", color=CARD), right=Side(style="thin", color=CARD),
                                      top=Side(style="thin", color=CARD), bottom=Side(style="thin", color=CARD))
            rango_valores.append(f"{get_column_letter(n0 + 1)}{r + 1}:{get_column_letter(n0 + len(banda))}{r + filas}")
            r += filas + 2
        rango = " ".join(rango_valores)
        primero = rango_valores[0].split(":")[0]
        abs_rango = ",".join(f"${x.split(':')[0][0]}${x.split(':')[0][1:]}:${x.split(':')[1][0]}${x.split(':')[1][1:]}"
                             for x in rango_valores)
        ws.conditional_formatting.add(rango, ColorScaleRule(start_type="min", start_color=SEQ_LO,
                                                            mid_type="percentile", mid_value=50, mid_color=SEQ_MID,
                                                            end_type="max", end_color=SEQ_HI))
        ws.conditional_formatting.add(rango, FormulaRule(formula=[f"{primero}>=PERCENTILE(({abs_rango}),0.6)"],
                                                         font=Font(color="FFFFFF", name=FONT, size=9, bold=True)))

    def dibujar_tabla(self, t, fila, n0):
        ws = self.dash
        enc = [t["cat"], self.nombre_medida, "% del total"]
        spans = [(n0, n0 + 1), (n0 + 2, n0 + 4), (n0 + 5, n0 + 5)]
        for (a, b), h in zip(spans, enc):
            if b > a:
                ws.merge_cells(start_row=fila, start_column=a, end_row=fila, end_column=b)
            c = ws.cell(fila, a, h)
            c.font = _font(8, True, INK_2)
            c.border = Border(bottom=Side(style="thin", color=INK_3))
            c.alignment = Alignment(horizontal="left" if a == n0 else "right")
        for i in range(t["f1"] - t["f0"] + 1):
            r, rc = fila + 1 + i, t["f0"] + i
            for (a, b) in spans:
                if b > a:
                    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            c = ws.cell(r, spans[0][0], f"=Calculos!$A${rc}")
            c.font = _font(9)
            c.alignment = Alignment(shrink_to_fit=True)
            v = ws.cell(r, spans[1][0], f"=Calculos!$B${rc}")
            v.number_format = self.fmt
            v.font = _font(9)
            pc = ws.cell(r, spans[2][0], f"=Calculos!$C${rc}")
            pc.number_format = "0.0%"
            pc.font = _font(9, color=INK_2)
            for (a, b) in spans:
                ws.cell(r, a).border = Border(bottom=Side(style="hair", color=GRID))
        ultima = fila + 1 + (t["f1"] - t["f0"])
        ws.conditional_formatting.add(
            f"{get_column_letter(spans[1][0])}{fila + 1}:{get_column_letter(spans[1][0])}{ultima}",
            DataBarRule(start_type="num", start_value=0, end_type="max", color="86B6EF",
                        showValue=True))

    def formatear_calc(self):
        ws = self.calc
        ws.column_dimensions["A"].width = 30
        for j in range(2, 15):
            ws.column_dimensions[get_column_letter(j)].width = 15
        ws.sheet_properties.tabColor = "B4B2AC"
        ws["H1"] = ("Hoja de apoyo: cada valor es una fórmula sobre la hoja Datos. "
                    "Si cambias o corriges datos en 'Datos', el Dashboard se actualiza solo.")
        ws["H1"].font = _font(9, color=INK_2)

    def escribir_info(self):
        ws = self.info
        p = self.p
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 60
        ws["A1"] = "Cómo se interpretaron los datos"
        ws["A1"].font = _font(14, True)
        filas = [("Archivo de origen", self.origen), ("Registros analizados", self.n - 1),
                 ("Medida principal", p.medida or "(conteo de registros)"),
                 ("Medida secundaria (dispersión)", p.medida2 or "—"),
                 ("Columna de fecha", p.fecha or "—"), ("Categoría principal", p.cat1 or "—"),
                 ("Categoría secundaria", p.cat2 or "—"), ("Categoría de participación", p.cat_dona or "—")]
        for i, (a, b) in enumerate(filas, start=3):
            ws.cell(i, 1, a).font = _font(10, True)
            ws.cell(i, 2, b).font = _font(10)
        r = len(filas) + 5
        for j, h in enumerate(["Columna", "Rol detectado", "Valores únicos", "Uso"], start=1):
            c = ws.cell(r, j, h)
            c.font = _font(10, True, "FFFFFF")
            c.fill = PatternFill("solid", fgColor="2A78D6")
        roles = {"numero": "Numérica", "fecha": "Fecha", "categoria": "Categoría", "id": "Identificador", "texto": "Texto libre"}
        usos = {"id": "Se ignora en los cálculos (identificador)", "texto": "Se ignora (demasiados valores distintos)"}
        for c in p.df.columns:
            r += 1
            t = p.tipos[c]
            uso = usos.get(t, "")
            if c == p.medida:
                uso = "Medida principal: KPIs, tendencia, rankings"
            elif c == p.medida2:
                uso = "Eje X de la dispersión"
            elif c == p.fecha:
                uso = "Eje de tiempo"
            elif c in (p.cat1, p.cat2, p.cat_dona):
                uso = "Agrupación en rankings / participación / mapa de calor"
            for j, v in enumerate([c, roles[t], int(p.df[c].nunique()), uso], start=1):
                ws.cell(r, j, v).font = _font(10)
        r += 2
        ws.cell(r, 1, "¿No eligió la columna que querías?").font = _font(10, True)
        ws.cell(r + 1, 1, 'Vuelve a ejecutar:  python dashboard_excel.py archivo.xlsx --medida "Columna" '
                          '--categoria "Columna" --fecha "Columna"').font = _font(10, color=INK_2)


# --------------------------------------------------------------------------
# API pública
# --------------------------------------------------------------------------
def generar_dashboard(entrada, salida=None, hoja=None, medida=None, categoria=None, fecha=None,
                      titulo=None) -> Path:
    entrada = Path(entrada)
    if not entrada.exists():
        raise FileNotFoundError(entrada)
    raw, nombre_hoja = leer_archivo(entrada, hoja)
    df = preparar_tabla(raw)
    if df.empty or df.shape[1] == 0:
        raise ValueError("No se encontraron filas de datos debajo de los encabezados.")
    perfil = perfilar(df, medida=medida, categoria=categoria, fecha=fecha)
    if not any([perfil.medida, perfil.cat1, perfil.fecha]):
        raise ValueError("No se encontró ninguna columna numérica, de fecha o de categorías para graficar.")
    origen = entrada.name + (f" › {nombre_hoja}" if entrada.suffix.lower() not in (".csv", ".txt", ".tsv") else "")
    titulo = titulo or f"Dashboard · {entrada.stem.replace('_', ' ').replace('-', ' ').title()}"
    wb = Constructor(perfil, titulo, origen).construir()
    salida = Path(salida) if salida else entrada.with_name(f"{entrada.stem}_dashboard.xlsx")
    wb.save(salida)
    return salida


def _pedir_archivo() -> str | None:
    try:
        import tkinter as tk
        from tkinter import filedialog
    except ImportError:
        return None
    raiz = tk.Tk()
    raiz.withdraw()
    ruta = filedialog.askopenfilename(title="Elige el archivo de datos",
                                      filetypes=[("Datos", "*.xlsx *.xlsm *.xls *.csv *.txt"), ("Todos", "*.*")])
    raiz.destroy()
    return ruta or None


def main(argv=None):
    ap = argparse.ArgumentParser(description="Crea un dashboard de Excel a partir de cualquier archivo de datos.")
    ap.add_argument("archivo", nargs="?", help="Archivo .xlsx/.xlsm/.xls/.csv")
    ap.add_argument("-o", "--salida", help="Ruta del Excel de salida (por defecto <archivo>_dashboard.xlsx)")
    ap.add_argument("--hoja", help="Hoja a analizar (por defecto, la que tenga más datos)")
    ap.add_argument("--medida", help="Columna numérica principal a analizar")
    ap.add_argument("--categoria", help="Columna de categoría principal")
    ap.add_argument("--fecha", help="Columna de fecha")
    ap.add_argument("--titulo", help="Título del dashboard")
    a = ap.parse_args(argv)
    archivo = a.archivo or _pedir_archivo()
    if not archivo:
        ap.error("indica un archivo de datos")
    try:
        salida = generar_dashboard(archivo, a.salida, a.hoja, a.medida, a.categoria, a.fecha, a.titulo)
    except (ValueError, FileNotFoundError) as e:
        print(f"Error: {e}", file=sys.stderr)
        return 1
    print(f"Dashboard creado: {salida}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
