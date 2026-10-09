"""Pruebas rápidas: python -m pytest -q"""
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
from dashboard_excel import generar_dashboard, parse_numero  # noqa: E402


@pytest.mark.parametrize("texto,esperado", [
    ("1.234,56", 1234.56), ("$1,234.56", 1234.56), ("(500)", -500), ("1 000", 1000),
    ("12,5", 12.5), ("1,000", 1000), ("45%", 45), ("-3.2", -3.2), ("abc", np.nan), ("2024-01-05", np.nan),
])
def test_parse_numero(texto, esperado):
    r = parse_numero(texto)
    assert (np.isnan(r) and np.isnan(esperado)) or r == pytest.approx(esperado)


def _hojas(ruta):
    return load_workbook(ruta).sheetnames


def test_reporte_con_titulo_y_total(tmp_path):
    n = 60
    df = pd.DataFrame({"Fecha": pd.date_range("2024-01-01", periods=n, freq="7D"),
                       "Zona": ["A", "B", "C"] * 20, "Ventas": np.arange(n) * 10.0})
    ruta = tmp_path / "rep.xlsx"
    with pd.ExcelWriter(ruta) as w:
        pd.DataFrame([["Mi reporte"]]).to_excel(w, header=False, index=False)
        df.to_excel(w, startrow=2, index=False)
        pd.DataFrame([["Total", None, df.Ventas.sum()]]).to_excel(w, startrow=n + 3, header=False, index=False)
    salida = generar_dashboard(ruta, tmp_path / "out.xlsx")
    assert _hojas(salida) == ["Dashboard", "Datos", "Calculos", "Acerca de"]
    datos = load_workbook(salida)["Datos"]
    assert datos.max_row == n + 1                     # la fila "Total" se descartó
    assert len(load_workbook(salida)["Dashboard"]._charts) >= 4


def test_csv_sin_numeros(tmp_path):
    ruta = tmp_path / "c.csv"
    ruta.write_text("Ciudad;Canal\n" + "\n".join(f"C{i % 4};K{i % 2}" for i in range(50)), encoding="utf-8")
    salida = generar_dashboard(ruta, tmp_path / "o.xlsx")
    assert "Dashboard" in _hojas(salida)              # cuenta registros cuando no hay medida


def test_sin_datos_utiles(tmp_path):
    ruta = tmp_path / "x.csv"
    ruta.write_text("Nota\n" + "\n".join(f"texto único {i}" for i in range(100)), encoding="utf-8")
    with pytest.raises(ValueError):
        generar_dashboard(ruta, tmp_path / "o.xlsx")
